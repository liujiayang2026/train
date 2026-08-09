from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, time, timedelta
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd


PROJECT_DIR = Path(__file__).resolve().parents[2]
RESULTS_DIR = PROJECT_DIR / "question-1" / "results"
PROCESSED_DIR = PROJECT_DIR / "question-1" / "data" / "processed"
QA_DIR = PROJECT_DIR / "question-1" / "qa"
DOCS_DIR = PROJECT_DIR / "question-1" / "docs"
YEAR_SHEETS = [str(year) for year in range(2016, 2022)]
SECONDS_PER_DAY = 24 * 3600


@dataclass
class LinearModel:
    name: str
    feature_names: list[str]
    mean: np.ndarray
    scale: np.ndarray
    coef: np.ndarray
    sigma2: float
    smearing: float


def _norm(value):
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def _is_number(value) -> bool:
    return isinstance(value, (int, float, np.number)) and not isinstance(value, bool)


def _parse_datetime(year: int, month: int, day: int, value) -> datetime:
    base = datetime(int(year), int(month), int(day))
    if isinstance(value, datetime):
        hours, minutes = value.hour, value.minute
    elif isinstance(value, time):
        hours, minutes = value.hour, value.minute
    elif isinstance(value, (int, float)):
        total_minutes = int(round(float(value) * 24 * 60))
        hours, minutes = divmod(total_minutes, 60)
    else:
        parts = str(value).strip().split(":")
        hours = int(parts[0])
        minutes = int(parts[1]) if len(parts) > 1 else 0
    return base + timedelta(hours=hours, minutes=minutes)


def find_attachment1() -> Path:
    candidates: list[tuple[Path, int]] = []
    for path in PROJECT_DIR.rglob("*.xlsx"):
        if path.name.startswith("~$"):
            continue
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            score = sum(1 for name in YEAR_SHEETS if name in wb.sheetnames)
            wb.close()
        except Exception:
            score = 0
        if score:
            candidates.append((path, score))
    if not candidates:
        raise FileNotFoundError("Could not find attachment 1 workbook with 2016-2021 sheets.")
    candidates.sort(key=lambda item: (-item[1], str(item[0])))
    return candidates[0][0]


def load_hydro_data(path: Path) -> pd.DataFrame:
    rows: list[dict] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet_name in YEAR_SHEETS:
        ws = wb[sheet_name]
        current_year = current_month = current_day = None
        for row in ws.iter_rows(min_row=2, values_only=True):
            year, month, day, sample_time, water_level, flow, sediment = [_norm(v) for v in row[:7]]
            if year is not None:
                current_year = int(year)
            if month is not None:
                current_month = int(month)
            if day is not None:
                current_day = int(day)
            if sample_time is None or current_year is None or current_month is None or current_day is None:
                continue
            dt = _parse_datetime(current_year, current_month, current_day, sample_time)
            rows.append(
                {
                    "datetime": dt,
                    "sheet_year": int(sheet_name),
                    "water_level_m": float(water_level) if _is_number(water_level) else np.nan,
                    "flow_m3s": float(flow) if _is_number(flow) else np.nan,
                    "sediment_obs_kgm3": float(sediment) if _is_number(sediment) else np.nan,
                }
            )
    wb.close()

    df = pd.DataFrame(rows).sort_values("datetime")
    df = (
        df.groupby("datetime", as_index=False)
        .agg(
            sheet_year=("sheet_year", "min"),
            water_level_m=("water_level_m", "mean"),
            flow_m3s=("flow_m3s", "mean"),
            sediment_obs_kgm3=("sediment_obs_kgm3", "mean"),
        )
        .sort_values("datetime")
        .reset_index(drop=True)
    )
    df["year"] = df["datetime"].dt.year
    df.loc[(df["year"] == 2022) & (df["sheet_year"] == 2021), "year"] = 2021
    df["water_level_m"] = df["water_level_m"].interpolate(limit_direction="both")
    return df


def day_of_year_fraction(datetimes: pd.Series) -> np.ndarray:
    values = []
    for dt in datetimes:
        start = datetime(dt.year, 1, 1)
        end = datetime(dt.year + 1, 1, 1)
        values.append((dt - start).total_seconds() / (end - start).total_seconds())
    return np.asarray(values, dtype=float)


def build_features(df: pd.DataFrame, kind: str) -> tuple[np.ndarray, list[str]]:
    t_years = (df["datetime"] - pd.Timestamp("2016-01-01")).dt.total_seconds().to_numpy() / (
        365.25 * SECONDS_PER_DAY
    )
    doy = day_of_year_fraction(df["datetime"])
    angle = 2.0 * np.pi * doy
    log_flow = np.log(df["flow_m3s"].to_numpy(dtype=float))
    water_level = df["water_level_m"].to_numpy(dtype=float)

    columns: list[tuple[str, np.ndarray]] = []
    if kind in {"time", "combined"}:
        columns.extend(
            [
                ("trend_year", t_years),
                ("sin_annual", np.sin(angle)),
                ("cos_annual", np.cos(angle)),
                ("sin_semiannual", np.sin(2.0 * angle)),
                ("cos_semiannual", np.cos(2.0 * angle)),
            ]
        )
    if kind in {"hydro", "combined"}:
        columns.extend(
            [
                ("log_flow", log_flow),
                ("water_level", water_level),
                ("log_flow_sq", log_flow**2),
                ("water_level_sq", water_level**2),
                ("log_flow_x_water_level", log_flow * water_level),
            ]
        )
    if kind == "simple_hydro_time":
        columns.extend(
            [
                ("trend_year", t_years),
                ("sin_annual", np.sin(angle)),
                ("cos_annual", np.cos(angle)),
                ("log_flow", log_flow),
                ("water_level", water_level),
            ]
        )

    return np.column_stack([values for _, values in columns]), [name for name, _ in columns]


def fit_linear_model(name: str, x: np.ndarray, y_log: np.ndarray, feature_names: list[str]) -> LinearModel:
    mean = np.nanmean(x, axis=0)
    scale = np.nanstd(x, axis=0)
    scale[scale == 0] = 1.0
    xz = (x - mean) / scale
    design = np.column_stack([np.ones(len(xz)), xz])
    coef = np.linalg.lstsq(design, y_log, rcond=None)[0]
    resid = y_log - design @ coef
    sigma2 = float(np.mean(resid**2))
    smearing = float(np.mean(np.exp(resid)))
    return LinearModel(name, feature_names, mean, scale, coef, sigma2, smearing)


def predict_log(model: LinearModel, x: np.ndarray) -> np.ndarray:
    xz = (x - model.mean) / model.scale
    design = np.column_stack([np.ones(len(xz)), xz])
    return design @ model.coef


def evaluate_predictions(y_true: np.ndarray, y_pred_log: np.ndarray) -> dict:
    y_log = np.log(y_true)
    resid_log = y_log - y_pred_log
    pred = np.exp(y_pred_log)
    resid = y_true - pred
    return {
        "rmse_log": float(np.sqrt(np.mean(resid_log**2))),
        "mae_log": float(np.mean(np.abs(resid_log))),
        "r2_log": float(1 - np.sum(resid_log**2) / np.sum((y_log - y_log.mean()) ** 2)),
        "rmse_kgm3": float(np.sqrt(np.mean(resid**2))),
        "mae_kgm3": float(np.mean(np.abs(resid))),
    }


def year_holdout_cv(train_df: pd.DataFrame, model_kinds: dict[str, str]) -> pd.DataFrame:
    records = []
    for model_name, kind in model_kinds.items():
        for year in sorted(train_df["year"].unique()):
            train_part = train_df[train_df["year"] != year]
            test_part = train_df[train_df["year"] == year]
            x_train, names = build_features(train_part, kind)
            y_train = np.log(train_part["sediment_obs_kgm3"].to_numpy(dtype=float))
            model = fit_linear_model(model_name, x_train, y_train, names)
            x_test, _ = build_features(test_part, kind)
            metrics = evaluate_predictions(
                test_part["sediment_obs_kgm3"].to_numpy(dtype=float),
                predict_log(model, x_test),
            )
            records.append({"model": model_name, "heldout_year": int(year), **metrics})
    cv = pd.DataFrame(records)
    summary = (
        cv.groupby("model", as_index=False)
        .agg(
            rmse_log=("rmse_log", "mean"),
            mae_log=("mae_log", "mean"),
            r2_log=("r2_log", "mean"),
            rmse_kgm3=("rmse_kgm3", "mean"),
            mae_kgm3=("mae_kgm3", "mean"),
        )
        .sort_values("rmse_log")
    )
    summary.insert(1, "heldout_year", "mean")
    return pd.concat([cv, summary], ignore_index=True)


def integrate_year(df_year: pd.DataFrame, value_col: str, year: int) -> float:
    part = df_year.sort_values("datetime").copy()
    start = pd.Timestamp(datetime(year, 1, 1))
    end = pd.Timestamp(datetime(year + 1, 1, 1))
    part = part[(part["datetime"] >= start) & (part["datetime"] <= end)]
    if part.empty:
        return float("nan")
    if part.iloc[0]["datetime"] > start:
        first = part.iloc[[0]].copy()
        first["datetime"] = start
        part = pd.concat([first, part], ignore_index=True)
    if part.iloc[-1]["datetime"] < end:
        last = part.iloc[[-1]].copy()
        last["datetime"] = end
        part = pd.concat([part, last], ignore_index=True)
    t_seconds = (part["datetime"] - start).dt.total_seconds().to_numpy(dtype=float)
    values = part[value_col].to_numpy(dtype=float)
    return float(np.trapezoid(values, t_seconds))


def summarize_relationships(train_df: pd.DataFrame, model: LinearModel, kind: str) -> pd.DataFrame:
    x, names = build_features(train_df, kind)
    coef_by_feature = {name: coef / scale for name, coef, scale in zip(names, model.coef[1:], model.scale)}
    observed = train_df.dropna(subset=["sediment_obs_kgm3", "flow_m3s", "water_level_m"]).copy()
    rows = [
        {
            "item": "corr_sediment_flow",
            "value": float(observed["sediment_obs_kgm3"].corr(observed["flow_m3s"])),
            "description": "Pearson correlation between observed sediment concentration and flow.",
        },
        {
            "item": "corr_log_sediment_log_flow",
            "value": float(np.corrcoef(np.log(observed["sediment_obs_kgm3"]), np.log(observed["flow_m3s"]))[0, 1]),
            "description": "Pearson correlation after log transform.",
        },
        {
            "item": "corr_sediment_water_level",
            "value": float(observed["sediment_obs_kgm3"].corr(observed["water_level_m"])),
            "description": "Pearson correlation between observed sediment concentration and water level.",
        },
    ]
    for name in names:
        rows.append(
            {
                "item": f"coef_{name}",
                "value": float(coef_by_feature[name]),
                "description": f"Approximate raw-scale coefficient of {name} in the log-concentration model.",
            }
        )
    return pd.DataFrame(rows)


def dataframe_to_markdown(df: pd.DataFrame, floatfmt: str = ".6g") -> str:
    headers = list(df.columns)
    body: list[list[str]] = []
    for _, row in df.iterrows():
        values = []
        for value in row:
            if isinstance(value, (float, np.floating)):
                values.append(format(float(value), floatfmt))
            else:
                values.append(str(value))
        body.append(values)
    widths = [
        max(len(str(header)), *(len(row[index]) for row in body)) if body else len(str(header))
        for index, header in enumerate(headers)
    ]
    header_line = "| " + " | ".join(str(header).ljust(widths[index]) for index, header in enumerate(headers)) + " |"
    sep_line = "| " + " | ".join("-" * widths[index] for index in range(len(headers))) + " |"
    row_lines = [
        "| " + " | ".join(value.ljust(widths[index]) for index, value in enumerate(row)) + " |"
        for row in body
    ]
    return "\n".join([header_line, sep_line, *row_lines])


def write_markdown(
    attachment_path: Path,
    df: pd.DataFrame,
    validation: pd.DataFrame,
    annual: pd.DataFrame,
    relationships: pd.DataFrame,
    monthly: pd.DataFrame,
    selected_model_name: str,
    selected_feature_names: list[str],
) -> None:
    best = validation[validation["heldout_year"].eq("mean")].sort_values("rmse_log").iloc[0]
    lines = [
        "# Question 1 Analysis",
        "",
        f"Input workbook: `{attachment_path}`",
        "",
        "## Data Cleaning",
        "",
        "- Parsed the 2016-2021 sheets, forward-filled year/month/day cells, and converted `24:00` records to the next day.",
        "- Water level and flow are nearly complete; missing water levels were time-interpolated before modeling.",
        "- Sediment concentration is sparse, so missing concentration values were estimated with a log-linear regression model.",
        "",
        "## Concentration Model",
        "",
        f"The selected model is `{selected_model_name}` and fits `log(C)` with: {', '.join(selected_feature_names)}.",
        "A log transform is used because sediment concentration is positive and strongly right-skewed.",
        "",
        f"Best mean held-out-year RMSE on log concentration: `{best['rmse_log']:.4f}` for `{best['model']}`.",
        "",
        "Key relationship diagnostics:",
        "",
        dataframe_to_markdown(relationships, floatfmt=".6g"),
        "",
        "Monthly time-pattern summary:",
        "",
        dataframe_to_markdown(
            monthly[
                [
                    "month",
                    "records",
                    "observed_sediment_records",
                    "mean_flow_m3s",
                    "mean_observed_sediment_kgm3",
                    "mean_used_sediment_kgm3",
                ]
            ],
            floatfmt=".6g",
        ),
        "",
        "## Annual Estimates",
        "",
        dataframe_to_markdown(
            annual[
                [
                    "year",
                    "records",
                    "sediment_observed_records",
                    "water_volume_1e8_m3",
                    "sediment_mass_1e4_t",
                    "mean_flow_m3s",
                    "mean_sediment_used_kgm3",
                ]
            ],
            floatfmt=".6g",
        ),
        "",
        "## Output Files",
        "",
        "- `data/processed/cleaned_hydro_timeseries.csv`: cleaned series with observed/model/used concentration.",
        "- `qa/model_validation.csv`: held-out-year model comparison.",
        "- `qa/relationship_diagnostics.csv`: correlation and model coefficient diagnostics.",
        "- `results/monthly_relationship_summary.csv`: month-level time relationship diagnostics.",
        "- `results/annual_flux_estimates.csv`: annual total water volume and sediment load estimates.",
        "",
    ]
    (DOCS_DIR / "question1-analysis.md").write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    QA_DIR.mkdir(parents=True, exist_ok=True)
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    attachment1 = find_attachment1()
    df = load_hydro_data(attachment1)

    train = df.dropna(subset=["flow_m3s", "water_level_m", "sediment_obs_kgm3"]).copy()
    train = train[train["sediment_obs_kgm3"] > 0].copy()
    model_kinds = {
        "time_only": "time",
        "hydro_only_quadratic": "hydro",
        "hydro_time_harmonic": "combined",
        "simple_hydro_time": "simple_hydro_time",
    }
    validation = year_holdout_cv(train, model_kinds)
    validation.to_csv(QA_DIR / "model_validation.csv", index=False, encoding="utf-8-sig")

    selected_model_name = str(
        validation[validation["heldout_year"].eq("mean")].sort_values("rmse_log").iloc[0]["model"]
    )
    selected_kind = model_kinds[selected_model_name]
    x_train, names = build_features(train, selected_kind)
    final_model = fit_linear_model(
        selected_model_name,
        x_train,
        np.log(train["sediment_obs_kgm3"].to_numpy(dtype=float)),
        names,
    )
    x_all, _ = build_features(df, selected_kind)
    pred_log = predict_log(final_model, x_all)
    df["sediment_model_kgm3"] = np.exp(pred_log) * final_model.smearing
    lower = max(0.001, float(train["sediment_obs_kgm3"].quantile(0.001)) * 0.5)
    upper = float(train["sediment_obs_kgm3"].quantile(0.999)) * 1.5
    df["sediment_model_kgm3"] = df["sediment_model_kgm3"].clip(lower=lower, upper=upper)
    df["sediment_filled_flag"] = df["sediment_obs_kgm3"].isna()
    df["sediment_used_kgm3"] = df["sediment_obs_kgm3"].fillna(df["sediment_model_kgm3"])
    df["sediment_flux_kg_s"] = df["flow_m3s"] * df["sediment_used_kgm3"]

    annual_records = []
    for year in range(2016, 2022):
        part = df[df["year"] == year].copy()
        water_volume_m3 = integrate_year(part, "flow_m3s", year)
        sediment_mass_kg = integrate_year(part, "sediment_flux_kg_s", year)
        seconds = (datetime(year + 1, 1, 1) - datetime(year, 1, 1)).total_seconds()
        annual_records.append(
            {
                "year": year,
                "records": int(len(part)),
                "sediment_observed_records": int(part["sediment_obs_kgm3"].notna().sum()),
                "sediment_observed_ratio": float(part["sediment_obs_kgm3"].notna().mean()),
                "water_volume_m3": water_volume_m3,
                "water_volume_1e8_m3": water_volume_m3 / 1e8,
                "sediment_mass_kg": sediment_mass_kg,
                "sediment_mass_t": sediment_mass_kg / 1000.0,
                "sediment_mass_1e4_t": sediment_mass_kg / 1e7,
                "mean_flow_m3s": water_volume_m3 / seconds,
                "mean_sediment_used_kgm3": float(part["sediment_used_kgm3"].mean()),
                "max_flow_m3s": float(part["flow_m3s"].max()),
                "max_sediment_used_kgm3": float(part["sediment_used_kgm3"].max()),
            }
        )
    annual = pd.DataFrame(annual_records)

    df["month"] = df["datetime"].dt.month
    monthly = (
        df[df["year"].between(2016, 2021)]
        .groupby("month", as_index=False)
        .agg(
            records=("datetime", "size"),
            observed_sediment_records=("sediment_obs_kgm3", lambda s: int(s.notna().sum())),
            mean_flow_m3s=("flow_m3s", "mean"),
            mean_water_level_m=("water_level_m", "mean"),
            mean_observed_sediment_kgm3=("sediment_obs_kgm3", "mean"),
            mean_used_sediment_kgm3=("sediment_used_kgm3", "mean"),
            mean_sediment_flux_kg_s=("sediment_flux_kg_s", "mean"),
        )
    )

    relationships = summarize_relationships(train, final_model, selected_kind)
    relationships.to_csv(QA_DIR / "relationship_diagnostics.csv", index=False, encoding="utf-8-sig")
    monthly.to_csv(RESULTS_DIR / "monthly_relationship_summary.csv", index=False, encoding="utf-8-sig")
    annual.to_csv(RESULTS_DIR / "annual_flux_estimates.csv", index=False, encoding="utf-8-sig")
    df.to_csv(PROCESSED_DIR / "cleaned_hydro_timeseries.csv", index=False, encoding="utf-8-sig")
    write_markdown(attachment1, df, validation, annual, relationships, monthly, selected_model_name, names)

    print("Attachment 1:", attachment1)
    print("Cleaned rows:", len(df))
    print("Training rows:", len(train))
    print("Best validation model:")
    print(validation[validation["heldout_year"].eq("mean")].sort_values("rmse_log").head(1).to_string(index=False))
    print("Selected final model:", selected_model_name)
    print("Annual estimates:")
    print(
        annual[
            [
                "year",
                "water_volume_1e8_m3",
                "sediment_mass_1e4_t",
                "mean_flow_m3s",
                "mean_sediment_used_kgm3",
            ]
        ].to_string(index=False)
    )


if __name__ == "__main__":
    main()
