from __future__ import annotations

import argparse
import platform
from dataclasses import dataclass
from datetime import datetime, time, timedelta
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd


ROUTE_DIR = Path(__file__).resolve().parents[1]
PACKAGE_DIR = Path(__file__).resolve().parents[5]
DEFAULT_INPUT = PACKAGE_DIR / "source" / "attachments" / "附件1.xlsx"
RUNS_DIR = ROUTE_DIR / "runs"
YEAR_SHEETS = [str(year) for year in range(2016, 2022)]
YEARS = list(range(2016, 2022))
SECONDS_PER_DAY = 24 * 3600
CANONICAL_BLOCK_HOURS = 72
REQUIRED_BLOCK_SENSITIVITY_HOURS = (24, 72, 168)
MODEL_KINDS = {
    "time_only": "time",
    "hydro_only_quadratic": "hydro",
    "hydro_time_harmonic": "combined",
    "simple_hydro_time": "simple_hydro_time",
}


@dataclass
class LinearModel:
    name: str
    feature_names: list[str]
    mean: np.ndarray
    scale: np.ndarray
    coef: np.ndarray
    sigma2: float
    smearing: float


@dataclass
class ResidualTimeBlockPlan:
    size: int
    target_groups: dict[int, list[tuple[np.ndarray, np.ndarray]]]
    source_positions: dict[int, list[np.ndarray]]
    source_hours: dict[int, list[np.ndarray]]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Run the canonical Q1 log-linear sediment analysis and validation."
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=DEFAULT_INPUT,
        help="Attachment 1 workbook (default: package source/attachments/附件1.xlsx).",
    )
    parser.add_argument(
        "--run-dir",
        type=Path,
        required=True,
        help="A new run-* directory under this route's runs directory.",
    )
    parser.add_argument(
        "--bootstrap-replicates",
        type=int,
        default=500,
        help="Number of within-year moving-time-block bootstrap draws per block length.",
    )
    parser.add_argument("--seed", type=int, default=20260809, help="Bootstrap random seed.")
    parser.add_argument(
        "--canonical-block-hours",
        type=int,
        default=CANONICAL_BLOCK_HOURS,
        help="Pre-specified canonical moving-block duration in hours (default: 72).",
    )
    parser.add_argument(
        "--block-sensitivity-hours",
        type=int,
        nargs="+",
        default=list(REQUIRED_BLOCK_SENSITIVITY_HOURS),
        help="Moving-block durations to compare; must include 24, 72, and 168 hours.",
    )
    return parser.parse_args()


def resolve_cli_path(path: Path) -> Path:
    return path.resolve() if path.is_absolute() else (Path.cwd() / path).resolve()


def validate_output_location(run_dir: Path) -> None:
    runs_root = RUNS_DIR.resolve()
    if run_dir.parent != runs_root:
        raise ValueError(f"--run-dir must be a direct child of {runs_root}")
    if not run_dir.name.startswith("run-"):
        raise ValueError("--run-dir basename must start with 'run-'.")
    if run_dir.exists() and not run_dir.is_dir():
        raise FileExistsError(f"--run-dir exists and is not a directory: {run_dir}")
    # A run record may be created before execution, as required by AGENTS.md.  No
    # other pre-existing item is safe to treat as an empty evidence directory.
    occupied = (
        sorted(path for path in run_dir.iterdir() if path.name != "run.md")
        if run_dir.exists()
        else []
    )
    if occupied:
        raise FileExistsError(
            "Refusing to overwrite generated evidence in an existing run: "
            + ", ".join(str(path) for path in occupied)
        )


def validate_block_hours(canonical: int, sensitivity: list[int]) -> list[int]:
    if canonical <= 0 or any(value <= 0 for value in sensitivity):
        raise ValueError("Moving-block durations must be positive hours.")
    values = list(dict.fromkeys(sensitivity))
    missing = sorted(set(REQUIRED_BLOCK_SENSITIVITY_HOURS) - set(values))
    if missing:
        raise ValueError(
            "--block-sensitivity-hours must include 24, 72, and 168; missing "
            + ", ".join(str(value) for value in missing)
        )
    if canonical not in values:
        raise ValueError("The canonical block duration must also appear in the sensitivity list.")
    return values


def report_input_path(path: Path) -> str:
    """Return a stable report label for package-local or external inputs."""
    resolved = path.resolve()
    try:
        return resolved.relative_to(PACKAGE_DIR.resolve()).as_posix()
    except ValueError:
        return str(resolved)


def _norm(value):
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def _as_finite_float(value) -> float:
    """Parse real numeric cells, including legal numeric text, or return NaN."""
    value = _norm(value)
    if value is None or isinstance(value, bool):
        return float("nan")
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return float("nan")
    return parsed if np.isfinite(parsed) else float("nan")


def _parse_datetime(year: int, month: int, day: int, value) -> datetime:
    base = datetime(int(year), int(month), int(day))
    if isinstance(value, datetime):
        hours, minutes = value.hour, value.minute
    elif isinstance(value, time):
        hours, minutes = value.hour, value.minute
    elif isinstance(value, (int, float)):
        total_minutes = int(round(float(value) * 24 * 60))
        hours, minutes = divmod(total_minutes, 60)
    else:
        parts = str(value).strip().split(":")
        hours = int(parts[0])
        minutes = int(parts[1]) if len(parts) > 1 else 0
    return base + timedelta(hours=hours, minutes=minutes)


def load_hydro_data(path: Path) -> pd.DataFrame:
    rows: list[dict] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    missing_sheets = [name for name in YEAR_SHEETS if name not in wb.sheetnames]
    if missing_sheets:
        wb.close()
        raise ValueError(f"Input workbook is missing year sheets: {missing_sheets}")

    for sheet_name in YEAR_SHEETS:
        ws = wb[sheet_name]
        current_year = current_month = current_day = None
        for row in ws.iter_rows(min_row=2, values_only=True):
            year, month, day, sample_time, water_level, flow, sediment = [_norm(v) for v in row[:7]]
            if year is not None:
                current_year = int(year)
            if month is not None:
                current_month = int(month)
            if day is not None:
                current_day = int(day)
            if sample_time is None or current_year is None or current_month is None or current_day is None:
                continue
            dt = _parse_datetime(current_year, current_month, current_day, sample_time)
            rows.append(
                {
                    "datetime": dt,
                    "sheet_year": int(sheet_name),
                    "water_level_m": _as_finite_float(water_level),
                    "flow_m3s": _as_finite_float(flow),
                    "sediment_obs_kgm3": _as_finite_float(sediment),
                }
            )
    wb.close()

    df = pd.DataFrame(rows).sort_values("datetime")
    df = (
        df.groupby("datetime", as_index=False)
        .agg(
            sheet_year=("sheet_year", "min"),
            water_level_m=("water_level_m", "mean"),
            flow_m3s=("flow_m3s", "mean"),
            sediment_obs_kgm3=("sediment_obs_kgm3", "mean"),
        )
        .sort_values("datetime")
        .reset_index(drop=True)
    )
    df["year"] = df["datetime"].dt.year
    # A 2021-sheet 24:00 record is the right endpoint of the 2021 integral.
    df.loc[(df["year"] == 2022) & (df["sheet_year"] == 2021), "year"] = 2021
    df["water_level_m"] = df["water_level_m"].interpolate(limit_direction="both")

    if df["flow_m3s"].isna().any():
        raise ValueError("Flow contains missing values; this route does not define flow imputation.")
    if (df["flow_m3s"] <= 0).any():
        raise ValueError("Flow must be positive because the model uses log(flow).")
    if df["water_level_m"].isna().any():
        raise ValueError("Water level remains missing after interpolation.")
    return df


def day_of_year_fraction(datetimes: pd.Series) -> np.ndarray:
    values = []
    for dt in datetimes:
        start = datetime(dt.year, 1, 1)
        end = datetime(dt.year + 1, 1, 1)
        values.append((dt - start).total_seconds() / (end - start).total_seconds())
    return np.asarray(values, dtype=float)


def build_features(df: pd.DataFrame, kind: str) -> tuple[np.ndarray, list[str]]:
    t_years = (df["datetime"] - pd.Timestamp("2016-01-01")).dt.total_seconds().to_numpy() / (
        365.25 * SECONDS_PER_DAY
    )
    doy = day_of_year_fraction(df["datetime"])
    angle = 2.0 * np.pi * doy
    log_flow = np.log(df["flow_m3s"].to_numpy(dtype=float))
    water_level = df["water_level_m"].to_numpy(dtype=float)

    columns: list[tuple[str, np.ndarray]] = []
    if kind in {"time", "combined"}:
        columns.extend(
            [
                ("trend_year", t_years),
                ("sin_annual", np.sin(angle)),
                ("cos_annual", np.cos(angle)),
                ("sin_semiannual", np.sin(2.0 * angle)),
                ("cos_semiannual", np.cos(2.0 * angle)),
            ]
        )
    if kind in {"hydro", "combined"}:
        columns.extend(
            [
                ("log_flow", log_flow),
                ("water_level", water_level),
                ("log_flow_sq", log_flow**2),
                ("water_level_sq", water_level**2),
                ("log_flow_x_water_level", log_flow * water_level),
            ]
        )
    if kind == "simple_hydro_time":
        columns.extend(
            [
                ("trend_year", t_years),
                ("sin_annual", np.sin(angle)),
                ("cos_annual", np.cos(angle)),
                ("log_flow", log_flow),
                ("water_level", water_level),
            ]
        )
    return np.column_stack([values for _, values in columns]), [name for name, _ in columns]


def fit_linear_model(name: str, x: np.ndarray, y_log: np.ndarray, feature_names: list[str]) -> LinearModel:
    mean = np.nanmean(x, axis=0)
    scale = np.nanstd(x, axis=0)
    scale[scale == 0] = 1.0
    xz = (x - mean) / scale
    design = np.column_stack([np.ones(len(xz)), xz])
    coef = np.linalg.lstsq(design, y_log, rcond=None)[0]
    resid = y_log - design @ coef
    sigma2 = float(np.mean(resid**2))
    smearing = float(np.mean(np.exp(resid)))
    return LinearModel(name, feature_names, mean, scale, coef, sigma2, smearing)


def predict_log(model: LinearModel, x: np.ndarray) -> np.ndarray:
    xz = (x - model.mean) / model.scale
    design = np.column_stack([np.ones(len(xz)), xz])
    return design @ model.coef


def model_bounds(
    train_df: pd.DataFrame,
    lower_quantile: float = 0.001,
    lower_multiplier: float = 0.5,
    upper_quantile: float = 0.999,
    upper_multiplier: float = 1.5,
) -> tuple[float, float]:
    observed = train_df["sediment_obs_kgm3"]
    lower = max(0.001, float(observed.quantile(lower_quantile)) * lower_multiplier)
    upper = float(observed.quantile(upper_quantile)) * upper_multiplier
    return lower, upper


def predict_concentration(
    model: LinearModel,
    df: pd.DataFrame,
    kind: str,
    bounds: tuple[float, float] | None,
) -> np.ndarray:
    x, _ = build_features(df, kind)
    predicted = np.exp(np.clip(predict_log(model, x), -50.0, 50.0)) * model.smearing
    if bounds is not None:
        predicted = np.clip(predicted, bounds[0], bounds[1])
    return predicted


def evaluate_predictions(y_true: np.ndarray, y_pred_log: np.ndarray) -> dict:
    y_log = np.log(y_true)
    resid_log = y_log - y_pred_log
    pred = np.exp(y_pred_log)
    resid = y_true - pred
    denominator = np.sum((y_log - y_log.mean()) ** 2)
    return {
        "rmse_log": float(np.sqrt(np.mean(resid_log**2))),
        "mae_log": float(np.mean(np.abs(resid_log))),
        "r2_log": float(1 - np.sum(resid_log**2) / denominator),
        "rmse_kgm3": float(np.sqrt(np.mean(resid**2))),
        "mae_kgm3": float(np.mean(np.abs(resid))),
    }


def year_holdout_cv(train_df: pd.DataFrame, model_kinds: dict[str, str]) -> pd.DataFrame:
    records = []
    for model_name, kind in model_kinds.items():
        for year in sorted(train_df["year"].unique()):
            train_part = train_df[train_df["year"] != year]
            test_part = train_df[train_df["year"] == year]
            x_train, names = build_features(train_part, kind)
            model = fit_linear_model(
                model_name,
                x_train,
                np.log(train_part["sediment_obs_kgm3"].to_numpy(dtype=float)),
                names,
            )
            x_test, _ = build_features(test_part, kind)
            metrics = evaluate_predictions(
                test_part["sediment_obs_kgm3"].to_numpy(dtype=float),
                predict_log(model, x_test),
            )
            records.append({"model": model_name, "heldout_year": int(year), **metrics})
    cv = pd.DataFrame(records)
    summary = (
        cv.groupby("model", as_index=False)
        .agg(
            rmse_log=("rmse_log", "mean"),
            mae_log=("mae_log", "mean"),
            r2_log=("r2_log", "mean"),
            rmse_kgm3=("rmse_kgm3", "mean"),
            mae_kgm3=("mae_kgm3", "mean"),
        )
        .sort_values("rmse_log")
    )
    summary.insert(1, "heldout_year", "mean")
    return pd.concat([cv, summary], ignore_index=True)


def trapezoid(values: np.ndarray, x: np.ndarray, axis: int = -1) -> np.ndarray | float:
    """Use NumPy 2.x trapezoid when available and NumPy 1.x trapz otherwise."""
    if hasattr(np, "trapezoid"):
        return np.trapezoid(values, x=x, axis=axis)
    return np.trapz(values, x=x, axis=axis)


def integrate_on_year_grid(datetimes: pd.Series, values: np.ndarray, year: int) -> float:
    order = np.argsort(datetimes.to_numpy())
    times = pd.to_datetime(datetimes.iloc[order]).reset_index(drop=True)
    y = np.asarray(values, dtype=float)[order]
    start = pd.Timestamp(datetime(year, 1, 1))
    end = pd.Timestamp(datetime(year + 1, 1, 1))
    keep = (times >= start) & (times <= end)
    times = times[keep].reset_index(drop=True)
    y = y[np.asarray(keep)]
    if len(times) == 0:
        return float("nan")
    if times.iloc[0] > start:
        times = pd.concat([pd.Series([start]), times], ignore_index=True)
        y = np.concatenate([[y[0]], y])
    if times.iloc[-1] < end:
        times = pd.concat([times, pd.Series([end])], ignore_index=True)
        y = np.concatenate([y, [y[-1]]])
    seconds = (times - start).dt.total_seconds().to_numpy(dtype=float)
    return float(trapezoid(y, seconds))


def integrate_year(df_year: pd.DataFrame, value_col: str, year: int) -> float:
    return integrate_on_year_grid(df_year["datetime"], df_year[value_col].to_numpy(dtype=float), year)


def build_completed_series(
    df: pd.DataFrame,
    model: LinearModel,
    kind: str,
    bounds: tuple[float, float] | None,
) -> pd.DataFrame:
    completed = df.copy()
    completed["sediment_model_kgm3"] = predict_concentration(model, completed, kind, bounds)
    completed["sediment_filled_flag"] = completed["sediment_obs_kgm3"].isna()
    completed["sediment_used_kgm3"] = completed["sediment_obs_kgm3"].fillna(
        completed["sediment_model_kgm3"]
    )
    completed["sediment_flux_kg_s"] = completed["flow_m3s"] * completed["sediment_used_kgm3"]
    return completed


def summarize_annual(df: pd.DataFrame) -> pd.DataFrame:
    records = []
    for year in YEARS:
        part = df[df["year"] == year].copy()
        water_volume_m3 = integrate_year(part, "flow_m3s", year)
        sediment_mass_kg = integrate_year(part, "sediment_flux_kg_s", year)
        seconds = (datetime(year + 1, 1, 1) - datetime(year, 1, 1)).total_seconds()
        records.append(
            {
                "year": year,
                "records": int(len(part)),
                "sediment_observed_records": int(part["sediment_obs_kgm3"].notna().sum()),
                "sediment_observed_ratio": float(part["sediment_obs_kgm3"].notna().mean()),
                "water_volume_m3": water_volume_m3,
                "water_volume_1e8_m3": water_volume_m3 / 1e8,
                "sediment_mass_kg": sediment_mass_kg,
                "sediment_mass_t": sediment_mass_kg / 1000.0,
                "sediment_mass_1e4_t": sediment_mass_kg / 1e7,
                "mean_flow_m3s": water_volume_m3 / seconds,
                "mean_sediment_used_kgm3": float(part["sediment_used_kgm3"].mean()),
                "max_flow_m3s": float(part["flow_m3s"].max()),
                "max_sediment_used_kgm3": float(part["sediment_used_kgm3"].max()),
            }
        )
    return pd.DataFrame(records)


def summarize_monthly(df: pd.DataFrame) -> pd.DataFrame:
    work = df.copy()
    work["month"] = work["datetime"].dt.month
    return (
        work[work["year"].between(2016, 2021)]
        .groupby("month", as_index=False)
        .agg(
            records=("datetime", "size"),
            observed_sediment_records=("sediment_obs_kgm3", lambda s: int(s.notna().sum())),
            mean_flow_m3s=("flow_m3s", "mean"),
            mean_water_level_m=("water_level_m", "mean"),
            mean_observed_sediment_kgm3=("sediment_obs_kgm3", "mean"),
            mean_used_sediment_kgm3=("sediment_used_kgm3", "mean"),
            mean_sediment_flux_kg_s=("sediment_flux_kg_s", "mean"),
        )
    )


def summarize_relationships(train_df: pd.DataFrame, model: LinearModel, kind: str) -> pd.DataFrame:
    _, names = build_features(train_df, kind)
    coef_by_feature = {
        name: coef / scale for name, coef, scale in zip(names, model.coef[1:], model.scale)
    }
    observed = train_df.dropna(
        subset=["sediment_obs_kgm3", "flow_m3s", "water_level_m"]
    ).copy()
    rows = [
        {
            "item": "observed_sample_size",
            "value": float(len(observed)),
            "description": "Number of positive observed-concentration records used for descriptive correlations.",
        },
        {
            "item": "corr_sediment_flow",
            "value": float(observed["sediment_obs_kgm3"].corr(observed["flow_m3s"])),
            "description": "Descriptive Pearson correlation between observed concentration and flow; not a significance test.",
        },
        {
            "item": "corr_log_sediment_log_flow",
            "value": float(
                np.corrcoef(
                    np.log(observed["sediment_obs_kgm3"]), np.log(observed["flow_m3s"])
                )[0, 1]
            ),
            "description": "Descriptive Pearson correlation after log transform; not a significance test.",
        },
        {
            "item": "corr_sediment_water_level",
            "value": float(observed["sediment_obs_kgm3"].corr(observed["water_level_m"])),
            "description": "Descriptive Pearson correlation between observed concentration and water level; Q/H collinearity is not removed.",
        },
    ]
    for name in names:
        rows.append(
            {
                "item": f"coef_{name}",
                "value": float(coef_by_feature[name]),
                "description": f"Raw-feature-scale coefficient of {name}; interpret jointly because polynomial terms are collinear.",
            }
        )
    return pd.DataFrame(rows)


def annual_holdout_flux_validation(
    observed: pd.DataFrame, selected_model_name: str, selected_kind: str
) -> pd.DataFrame:
    records = []
    for year in YEARS:
        train_part = observed[observed["year"] != year].copy()
        test_part = observed[observed["year"] == year].copy().sort_values("datetime")
        x_train, names = build_features(train_part, selected_kind)
        model = fit_linear_model(
            selected_model_name,
            x_train,
            np.log(train_part["sediment_obs_kgm3"].to_numpy(dtype=float)),
            names,
        )
        predicted = predict_concentration(model, test_part, selected_kind, model_bounds(train_part))
        flow = test_part["flow_m3s"].to_numpy(dtype=float)
        observed_flux = flow * test_part["sediment_obs_kgm3"].to_numpy(dtype=float)
        predicted_flux = flow * predicted
        observed_mass = integrate_on_year_grid(test_part["datetime"], observed_flux, year)
        predicted_mass = integrate_on_year_grid(test_part["datetime"], predicted_flux, year)
        absolute_error_mass = integrate_on_year_grid(
            test_part["datetime"], np.abs(predicted_flux - observed_flux), year
        )
        records.append(
            {
                "heldout_year": year,
                "observed_records": int(len(test_part)),
                "grid_start": test_part["datetime"].min(),
                "grid_end": test_part["datetime"].max(),
                "observed_grid_mass_1e4_t": observed_mass / 1e7,
                "predicted_grid_mass_1e4_t": predicted_mass / 1e7,
                "signed_error_pct": 100.0 * (predicted_mass - observed_mass) / observed_mass,
                "flux_wape_pct": 100.0 * absolute_error_mass / observed_mass,
            }
        )
    return pd.DataFrame(records)


def annual_mass_map(df: pd.DataFrame) -> dict[int, float]:
    return {
        year: integrate_year(df[df["year"] == year], "sediment_flux_kg_s", year) / 1e7
        for year in YEARS
    }


def build_sensitivity(
    df: pd.DataFrame,
    observed: pd.DataFrame,
    baseline: pd.DataFrame,
    selected_model_name: str,
    selected_kind: str,
    final_model: LinearModel,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    baseline_mass = dict(zip(baseline["year"], baseline["sediment_mass_1e4_t"]))
    records: list[dict] = []

    def add_scenario(group: str, scenario: str, description: str, mass_by_year: dict[int, float]) -> None:
        for year in YEARS:
            value = mass_by_year[year]
            reference = float(baseline_mass[year])
            records.append(
                {
                    "scenario_group": group,
                    "scenario": scenario,
                    "description": description,
                    "year": year,
                    "sediment_mass_1e4_t": value,
                    "relative_change_pct_vs_baseline": 100.0 * (value - reference) / reference,
                }
            )

    add_scenario(
        "reference",
        "full_fit_selected_standard_clip",
        "Selected four-model CV winner, all observed years fitted, standard clipping.",
        {year: float(baseline_mass[year]) for year in YEARS},
    )

    standard_bounds = model_bounds(observed)
    for model_name, kind in MODEL_KINDS.items():
        if model_name == selected_model_name:
            continue
        x_train, names = build_features(observed, kind)
        model = fit_linear_model(
            model_name,
            x_train,
            np.log(observed["sediment_obs_kgm3"].to_numpy(dtype=float)),
            names,
        )
        completed = build_completed_series(df, model, kind, standard_bounds)
        add_scenario(
            "candidate_model",
            f"model_{model_name}",
            "Alternative among the four machine-evidenced candidate models; standard clipping.",
            annual_mass_map(completed),
        )

    clipping_scenarios: list[tuple[str, str, tuple[float, float] | None]] = [
        ("clip_none", "No prediction clipping.", None),
        (
            "clip_tight_q01_q99",
            "Clip to 0.5*q01 and 1.25*q99 of observed concentration.",
            model_bounds(observed, 0.01, 0.5, 0.99, 1.25),
        ),
        (
            "clip_loose_q0005_q9995",
            "Clip to 0.25*q0005 and 2*q9995 of observed concentration.",
            model_bounds(observed, 0.0005, 0.25, 0.9995, 2.0),
        ),
    ]
    for scenario, description, bounds in clipping_scenarios:
        completed = build_completed_series(df, final_model, selected_kind, bounds)
        add_scenario("prediction_clip", scenario, description, annual_mass_map(completed))

    crossfit_mass: dict[int, float] = {}
    for year in YEARS:
        train_part = observed[observed["year"] != year].copy()
        target = df[df["year"] == year].copy()
        x_train, names = build_features(train_part, selected_kind)
        model = fit_linear_model(
            selected_model_name,
            x_train,
            np.log(train_part["sediment_obs_kgm3"].to_numpy(dtype=float)),
            names,
        )
        completed = build_completed_series(target, model, selected_kind, model_bounds(train_part))
        crossfit_mass[year] = integrate_year(completed, "sediment_flux_kg_s", year) / 1e7
    add_scenario(
        "training_window",
        "leave_target_year_out",
        "Each target year's missing concentrations are filled by a model fitted without that year's observations.",
        crossfit_mass,
    )

    sensitivity = pd.DataFrame(records)
    summary_rows = []
    for year in YEARS:
        part = sensitivity[sensitivity["year"] == year]
        row: dict[str, float | int] = {
            "year": year,
            "baseline_sediment_mass_1e4_t": float(baseline_mass[year]),
            "all_scenarios_min_1e4_t": float(part["sediment_mass_1e4_t"].min()),
            "all_scenarios_max_1e4_t": float(part["sediment_mass_1e4_t"].max()),
            "all_scenarios_max_abs_change_pct": float(
                part["relative_change_pct_vs_baseline"].abs().max()
            ),
        }
        for group in ("candidate_model", "prediction_clip", "training_window"):
            group_part = part[part["scenario_group"] == group]
            row[f"{group}_max_abs_change_pct"] = float(
                group_part["relative_change_pct_vs_baseline"].abs().max()
            )
        summary_rows.append(row)
    return sensitivity, pd.DataFrame(summary_rows)


def build_moving_time_block_pool(
    observed: pd.DataFrame, block_hours: int
) -> dict[int, list[np.ndarray]]:
    """Build non-circular, within-year blocks using actual elapsed time."""
    if block_hours <= 0:
        raise ValueError("Moving-block duration must be positive.")
    duration = pd.Timedelta(hours=block_hours)
    pools: dict[int, list[np.ndarray]] = {}
    years = observed["year"].to_numpy(dtype=int)
    for year in YEARS:
        positions = np.flatnonzero(years == year)
        if len(positions) == 0:
            raise ValueError(f"No observed concentration records are available for {year}.")
        times = pd.to_datetime(observed.iloc[positions]["datetime"]).reset_index(drop=True)
        if not times.is_monotonic_increasing:
            raise ValueError("Observed records must be sorted by real datetime before blocking.")
        year_end = pd.Timestamp(datetime(year + 1, 1, 1))
        valid_starts = times[times + duration <= year_end]
        if valid_starts.empty:
            raise ValueError(
                f"No complete {block_hours}-hour moving block is available within {year}."
            )
        blocks: list[np.ndarray] = []
        for start in valid_starts:
            in_window = (times >= start) & (times < start + duration)
            block = positions[np.flatnonzero(in_window.to_numpy())]
            if len(block):
                blocks.append(block)
        if not blocks:
            raise ValueError(f"No non-empty {block_hours}-hour blocks were built for {year}.")
        pools[year] = blocks
    return pools


def sample_pairs_from_time_blocks(
    pools: dict[int, list[np.ndarray]],
    observed: pd.DataFrame,
    rng: np.random.Generator,
) -> np.ndarray:
    """Sample complete chronological blocks within each year, never individual rows."""
    years = observed["year"].to_numpy(dtype=int)
    sampled: list[np.ndarray] = []
    for year in YEARS:
        required = int(np.sum(years == year))
        selected: list[np.ndarray] = []
        count = 0
        while count < required:
            block = pools[year][int(rng.integers(0, len(pools[year])))]
            selected.append(block)
            count += len(block)
        # Truncating only the tail of the last chronological block gives the
        # original per-year sample size without reverting to row resampling.
        sampled.append(np.concatenate(selected)[:required])
    return np.concatenate(sampled)


def _nearest_values(
    source_hours: np.ndarray, source_values: np.ndarray, target_hours: np.ndarray
) -> np.ndarray:
    right = np.searchsorted(source_hours, target_hours, side="left")
    right = np.clip(right, 0, len(source_hours) - 1)
    left = np.clip(right - 1, 0, len(source_hours) - 1)
    choose_right = np.abs(source_hours[right] - target_hours) < np.abs(
        source_hours[left] - target_hours
    )
    nearest = np.where(choose_right, right, left)
    return source_values[nearest]


def _datetime_ns(values: pd.Series) -> np.ndarray:
    """Normalize pandas 2/3 datetime storage units to explicit nanoseconds."""
    return pd.to_datetime(values).to_numpy(dtype="datetime64[ns]").astype(np.int64)


def build_residual_time_block_plan(
    df: pd.DataFrame,
    observed: pd.DataFrame,
    pools: dict[int, list[np.ndarray]],
    block_hours: int,
) -> ResidualTimeBlockPlan:
    """Cache real-time geometry so bootstrap draws only perform NumPy indexing."""
    missing = df["sediment_obs_kgm3"].isna().to_numpy()
    df_years = df["year"].to_numpy(dtype=int)
    duration_seconds = float(block_hours * 3600)
    observed_ns = _datetime_ns(observed["datetime"])
    source_positions: dict[int, list[np.ndarray]] = {}
    source_hours: dict[int, list[np.ndarray]] = {}
    target_groups: dict[int, list[tuple[np.ndarray, np.ndarray]]] = {}
    for year in YEARS:
        source_positions[year] = pools[year]
        source_hours[year] = [
            (observed_ns[positions] - observed_ns[positions[0]]) / 3.6e12
            for positions in pools[year]
        ]
        target_positions = np.flatnonzero(missing & (df_years == year))
        if len(target_positions) == 0:
            target_groups[year] = []
            continue
        target_ns = _datetime_ns(df.iloc[target_positions]["datetime"])
        year_start_ns = pd.Timestamp(datetime(year, 1, 1)).value
        elapsed_seconds = (target_ns - year_start_ns) / 1e9
        block_ids = np.floor(
            elapsed_seconds / duration_seconds
        ).astype(int)
        groups: list[tuple[np.ndarray, np.ndarray]] = []
        for block_id in np.unique(block_ids):
            local_target = np.flatnonzero(block_ids == block_id)
            target_hours = (
                elapsed_seconds[local_target] - int(block_id) * duration_seconds
            )
            groups.append((target_positions[local_target], target_hours / 3600.0))
        target_groups[year] = groups
    return ResidualTimeBlockPlan(
        size=len(df),
        target_groups=target_groups,
        source_positions=source_positions,
        source_hours=source_hours,
    )


def sample_missing_residual_time_path(
    plan: ResidualTimeBlockPlan,
    observed_residuals: np.ndarray,
    rng: np.random.Generator,
) -> np.ndarray:
    """Map sampled empirical residual blocks to cached equal-duration target blocks."""
    sampled = np.zeros(plan.size, dtype=float)
    for year in YEARS:
        for target_positions, target_hours in plan.target_groups[year]:
            source_index = int(rng.integers(0, len(plan.source_positions[year])))
            positions = plan.source_positions[year][source_index]
            sampled[target_positions] = _nearest_values(
                plan.source_hours[year][source_index],
                observed_residuals[positions],
                target_hours,
            )
    return sampled


def build_annual_integration_plan(
    df: pd.DataFrame,
) -> dict[int, tuple[np.ndarray, np.ndarray]]:
    """Precompute trapezoid weights for the unchanged full-series year grids."""
    plan: dict[int, tuple[np.ndarray, np.ndarray]] = {}
    years = df["year"].to_numpy(dtype=int)
    datetime_ns = _datetime_ns(df["datetime"])
    for year in YEARS:
        positions = np.flatnonzero(years == year)
        start_ns = pd.Timestamp(datetime(year, 1, 1)).value
        end_ns = pd.Timestamp(datetime(year + 1, 1, 1)).value
        keep = (datetime_ns[positions] >= start_ns) & (datetime_ns[positions] <= end_ns)
        positions = positions[keep]
        if len(positions) == 0:
            raise ValueError(f"No integration grid is available for {year}.")
        order = np.argsort(datetime_ns[positions])
        positions = positions[order]
        times = datetime_ns[positions]
        origins = np.arange(len(positions), dtype=int)
        if times[0] > start_ns:
            times = np.concatenate([[start_ns], times])
            origins = np.concatenate([[0], origins])
        if times[-1] < end_ns:
            times = np.concatenate([times, [end_ns]])
            origins = np.concatenate([origins, [len(positions) - 1]])
        seconds = (times - start_ns) / 1e9
        expanded_weights = np.zeros(len(seconds), dtype=float)
        if len(seconds) > 1:
            expanded_weights[0] = (seconds[1] - seconds[0]) / 2.0
            expanded_weights[-1] = (seconds[-1] - seconds[-2]) / 2.0
        if len(seconds) > 2:
            expanded_weights[1:-1] = (seconds[2:] - seconds[:-2]) / 2.0
        weights = np.bincount(
            origins, weights=expanded_weights, minlength=len(positions)
        ).astype(float)
        plan[year] = positions, weights
    return plan


def build_time_block_diagnostics(
    observed: pd.DataFrame, block_hours_values: list[int], canonical_block_hours: int
) -> pd.DataFrame:
    rows = []
    for block_hours in block_hours_values:
        pools = build_moving_time_block_pool(observed, block_hours)
        for year in YEARS:
            counts = np.asarray([len(block) for block in pools[year]], dtype=int)
            rows.append(
                {
                    "time_block_hours": block_hours,
                    "is_canonical": block_hours == canonical_block_hours,
                    "year": year,
                    "observed_records": int((observed["year"] == year).sum()),
                    "candidate_blocks": int(len(counts)),
                    "records_per_block_min": int(counts.min()),
                    "records_per_block_median": float(np.median(counts)),
                    "records_per_block_max": int(counts.max()),
                    "block_scope": "non-circular within-year real-time window",
                }
            )
    return pd.DataFrame(rows)


def bootstrap_annual_uncertainty(
    df: pd.DataFrame,
    observed: pd.DataFrame,
    baseline: pd.DataFrame,
    selected_model_name: str,
    selected_kind: str,
    replicates: int,
    seed: int,
    block_hours: int,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    if replicates < 100:
        raise ValueError("Use at least 100 bootstrap replicates for stable percentile summaries.")
    if block_hours < 1:
        raise ValueError("Moving-block duration must be positive.")

    rng = np.random.default_rng(seed)
    missing = df["sediment_obs_kgm3"].isna().to_numpy()
    observed_values = df["sediment_obs_kgm3"].to_numpy(dtype=float)
    flow = df["flow_m3s"].to_numpy(dtype=float)
    all_features, _ = build_features(df, selected_kind)
    observed_features, _ = build_features(observed, selected_kind)
    observed_log = np.log(observed["sediment_obs_kgm3"].to_numpy(dtype=float))
    block_pools = build_moving_time_block_pool(observed, block_hours)
    residual_plan = build_residual_time_block_plan(df, observed, block_pools, block_hours)
    integration_plan = build_annual_integration_plan(df)
    draws = np.empty((replicates, len(YEARS)), dtype=float)

    for draw in range(replicates):
        sampled_indices = sample_pairs_from_time_blocks(block_pools, observed, rng)
        boot = observed.iloc[sampled_indices].reset_index(drop=True)
        x_boot, names = build_features(boot, selected_kind)
        y_boot = np.log(boot["sediment_obs_kgm3"].to_numpy(dtype=float))
        model = fit_linear_model(selected_model_name, x_boot, y_boot, names)
        residuals = observed_log - predict_log(model, observed_features)
        residuals = residuals - float(residuals.mean())
        sampled_residuals = sample_missing_residual_time_path(residual_plan, residuals, rng)
        pred_log_all = predict_log(model, all_features)
        simulated_missing = np.exp(
            np.clip(pred_log_all[missing] + sampled_residuals[missing], -50.0, 50.0)
        )
        bounds = model_bounds(boot)
        simulated_missing = np.clip(simulated_missing, bounds[0], bounds[1])
        used = observed_values.copy()
        used[missing] = simulated_missing
        flux = flow * used
        for year_index, year in enumerate(YEARS):
            positions, weights = integration_plan[year]
            draws[draw, year_index] = float(np.dot(flux[positions], weights) / 1e7)

    draw_table = pd.DataFrame(
        {
            "draw": np.repeat(np.arange(1, replicates + 1), len(YEARS)),
            "year": np.tile(YEARS, replicates),
            "sediment_mass_1e4_t": draws.reshape(-1),
            "time_block_hours": block_hours,
        }
    )
    baseline_map = dict(zip(baseline["year"], baseline["sediment_mass_1e4_t"]))
    summary = []
    for year_index, year in enumerate(YEARS):
        values = draws[:, year_index]
        lower, median, upper = np.quantile(values, [0.025, 0.5, 0.975])
        base = float(baseline_map[year])
        summary.append(
            {
                "year": year,
                "baseline_sediment_mass_1e4_t": base,
                "bootstrap_median_1e4_t": float(median),
                "bootstrap_p2_5_1e4_t": float(lower),
                "bootstrap_p97_5_1e4_t": float(upper),
                "relative_half_width_pct": float(100.0 * (upper - lower) / (2.0 * base)),
                "baseline_inside_interval": bool(lower <= base <= upper),
                "bootstrap_replicates": replicates,
                "random_seed": seed,
                "time_block_hours": block_hours,
                "pairs_resampling": "within-year moving real-time blocks",
                "residual_resampling": "within-year moving real-time blocks mapped by relative time",
            }
        )
    return draw_table, pd.DataFrame(summary)


def dataframe_to_markdown(df: pd.DataFrame, floatfmt: str = ".6g") -> str:
    headers = list(df.columns)
    body: list[list[str]] = []
    for _, row in df.iterrows():
        values = []
        for value in row:
            if isinstance(value, (float, np.floating)):
                values.append(format(float(value), floatfmt))
            else:
                values.append(str(value))
        body.append(values)
    widths = [
        max(len(str(header)), *(len(row[index]) for row in body)) if body else len(str(header))
        for index, header in enumerate(headers)
    ]
    header_line = "| " + " | ".join(
        str(header).ljust(widths[index]) for index, header in enumerate(headers)
    ) + " |"
    sep_line = "| " + " | ".join("-" * widths[index] for index in range(len(headers))) + " |"
    row_lines = [
        "| " + " | ".join(value.ljust(widths[index]) for index, value in enumerate(row)) + " |"
        for row in body
    ]
    return "\n".join([header_line, sep_line, *row_lines])


def build_quality_checks(
    df: pd.DataFrame,
    annual: pd.DataFrame,
    uncertainty: pd.DataFrame,
    block_sensitivity: pd.DataFrame,
    block_diagnostics: pd.DataFrame,
    selected_model_name: str,
    canonical_block_hours: int,
) -> pd.DataFrame:
    sensitivity_hours = set(block_sensitivity["time_block_hours"].astype(int))
    checks = [
        ("rows_present", len(df) > 0, f"rows={len(df)}"),
        (
            "all_six_years_present",
            set(df["year"].unique()) == set(YEARS),
            f"years={sorted(df['year'].unique().tolist())}",
        ),
        ("unique_datetimes", df["datetime"].is_unique, f"duplicates={df['datetime'].duplicated().sum()}"),
        ("flow_complete_positive", bool(df["flow_m3s"].notna().all() and (df["flow_m3s"] > 0).all()), "required for log(flow) and integration"),
        ("water_level_complete", bool(df["water_level_m"].notna().all()), "after documented interpolation"),
        ("used_sediment_complete_positive", bool(df["sediment_used_kgm3"].notna().all() and (df["sediment_used_kgm3"] > 0).all()), "observed where available, modeled otherwise"),
        ("four_models_compared", len(MODEL_KINDS) == 4, f"models={','.join(MODEL_KINDS)}"),
        ("selected_model_is_cv_candidate", selected_model_name in MODEL_KINDS, selected_model_name),
        ("six_annual_estimates", len(annual) == 6 and annual["sediment_mass_kg"].gt(0).all(), "2016-2021 positive totals"),
        ("bootstrap_intervals_ordered", bool((uncertainty["bootstrap_p2_5_1e4_t"] <= uncertainty["bootstrap_median_1e4_t"]).all() and (uncertainty["bootstrap_median_1e4_t"] <= uncertainty["bootstrap_p97_5_1e4_t"]).all()), "p2.5 <= median <= p97.5"),
        (
            "canonical_block_is_prespecified",
            canonical_block_hours == CANONICAL_BLOCK_HOURS,
            f"canonical_hours={canonical_block_hours}; prespecified={CANONICAL_BLOCK_HOURS}",
        ),
        (
            "required_time_block_sensitivity",
            set(REQUIRED_BLOCK_SENSITIVITY_HOURS).issubset(sensitivity_hours),
            f"hours={sorted(sensitivity_hours)}",
        ),
        (
            "time_blocks_are_within_year",
            bool(
                len(block_diagnostics) == len(sensitivity_hours) * len(YEARS)
                and block_diagnostics["candidate_blocks"].gt(0).all()
                and block_diagnostics["block_scope"]
                .eq("non-circular within-year real-time window")
                .all()
            ),
            "non-circular real-time blocks; one diagnostic row per year and duration",
        ),
    ]
    return pd.DataFrame(
        [{"check": name, "passed": bool(passed), "detail": detail} for name, passed, detail in checks]
    )


def write_analysis(
    path: Path,
    input_path: Path,
    df: pd.DataFrame,
    validation: pd.DataFrame,
    annual: pd.DataFrame,
    holdout: pd.DataFrame,
    sensitivity_summary: pd.DataFrame,
    uncertainty: pd.DataFrame,
    block_sensitivity: pd.DataFrame,
    selected_model_name: str,
    selected_feature_names: list[str],
    canonical_block_hours: int,
) -> None:
    mean_cv = validation[validation["heldout_year"].eq("mean")].sort_values("rmse_log")
    observed_ratio = df["sediment_obs_kgm3"].notna().mean()
    annual_view = annual[
        ["year", "water_volume_1e8_m3", "sediment_mass_1e4_t", "sediment_observed_ratio"]
    ]
    holdout_view = holdout[
        ["heldout_year", "observed_records", "signed_error_pct", "flux_wape_pct"]
    ]
    uncertainty_view = uncertainty[
        [
            "year",
            "baseline_sediment_mass_1e4_t",
            "bootstrap_p2_5_1e4_t",
            "bootstrap_p97_5_1e4_t",
            "relative_half_width_pct",
        ]
    ]
    block_sensitivity_view = block_sensitivity[
        [
            "time_block_hours",
            "year",
            "bootstrap_p2_5_1e4_t",
            "bootstrap_p97_5_1e4_t",
            "relative_half_width_pct",
        ]
    ]
    lines = [
        "# 问题一 canonical 运行摘要",
        "",
        f"- 输入：`{report_input_path(input_path)}`",
        f"- 清洗后记录数：{len(df)}",
        f"- 含沙量实测比例：{observed_ratio:.2%}；模型补全比例：{1-observed_ratio:.2%}",
        f"- 本次四模型交叉验证胜者：`{selected_model_name}`",
        f"- 入模特征：{', '.join(selected_feature_names)}",
        "",
        "这里的“胜者”仅表示本次机器比较中 RMSE_log 最低，不代表人工 adopted 决定。",
        "",
        "## 四个候选模型的留一年平均验证",
        "",
        dataframe_to_markdown(mean_cv, ".6g"),
        "",
        "## 年度积分结果",
        "",
        dataframe_to_markdown(annual_view, ".6g"),
        "",
        "## 留一年年度输沙聚合验证",
        "",
        dataframe_to_markdown(holdout_view, ".6g"),
        "",
        "该验证只在留出年份的实测含沙量时间网格上比较 Q×C 的积分，不是完整年度真值。",
        "",
        "## 模型补全自助法区间",
        "",
        f"预先指定 {canonical_block_hours} 小时为主口径。系数样本和残差路径都按年份从真实时间轴上的连续移动块抽取，不逐行重采样。下表是当前模型内的经验百分位区间。",
        "",
        dataframe_to_markdown(uncertainty_view, ".6g"),
        "",
        "区间传播了移动时间块训练样本与残差路径对缺失含沙量补全的影响；不含流量测量误差、模型族选择误差和未观测极端事件结构偏差，不能表述为完整误差的正式置信区间。",
        "",
        "## 24/72/168小时块长敏感性",
        "",
        dataframe_to_markdown(block_sensitivity_view, ".6g"),
        "",
        "## 敏感性摘要",
        "",
        dataframe_to_markdown(sensitivity_summary, ".6g"),
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def write_validation_summary(
    path: Path,
    holdout: pd.DataFrame,
    sensitivity_summary: pd.DataFrame,
    uncertainty: pd.DataFrame,
    block_sensitivity: pd.DataFrame,
    block_diagnostics: pd.DataFrame,
    replicates: int,
    seed: int,
    canonical_block_hours: int,
    block_hours_values: list[int],
) -> None:
    lines = [
        "# Q1 年度排沙端到端验证与敏感性说明",
        "",
        "## 验证对象",
        "",
        "验证链为：实测水位/流量与稀疏含沙量 → 含沙量模型 → 缺失值补全 → Q×C → 梯形年度积分。水量积分不依赖含沙量模型。",
        "",
        "## 留一年年度聚合检验",
        "",
        "每次完全排除目标年份的含沙量观测后拟合，再在该年份所有实测含沙量时刻预测，并在同一稀疏时间网格上分别积分实测 Q×C 与预测 Q×C。这避免了目标年份浓度进入拟合，但稀疏网格不等于完整年度真值。",
        "",
        dataframe_to_markdown(holdout, ".6g"),
        "",
        "## 完整年度敏感性",
        "",
        "逐年比较四模型、预测截尾边界以及排除目标年份训练三类扰动。`all_scenarios_max_abs_change_pct` 是列入场景相对基准的最大绝对变化，并非概率区间。",
        "",
        dataframe_to_markdown(sensitivity_summary, ".6g"),
        "",
        "## 自助法不确定性",
        "",
        f"预先指定 {canonical_block_hours} 小时为主块长，每个块长各运行 {replicates} 次，随机种子 {seed}。系数的pairs bootstrap按年份从原始实测记录的真实时间轴抽取不跨年、非循环的连续移动块，按块拼接到原年份样本量；没有逐行独立重采样。每次重新拟合后，在原始实测时刻计算中心化log残差；目标年份按同样时长划分日历块，从同一年抽源残差块，并按块内相对时间最近邻映射到缺失时刻。",
        "",
        dataframe_to_markdown(uncertainty, ".6g"),
        "",
        "以上是当前模型内的经验百分位区间，不是覆盖所有误差源的置信区间。",
        "",
        "## 块长敏感性",
        "",
        f"预先要求比较 {', '.join(str(value) for value in block_hours_values)} 小时；72小时不是查看结果后选出的最窄区间。",
        "",
        dataframe_to_markdown(block_sensitivity, ".6g"),
        "",
        "## 时间块诊断",
        "",
        "每个候选块都由同一年真实时间窗口 `[start, start+B)` 内按时间排序的实测记录组成，不在年末循环回卷。记录数因原始采样不等间隔而自然变化。",
        "",
        dataframe_to_markdown(block_diagnostics, ".6g"),
        "",
        "## 解释边界",
        "",
        "- 自助区间只反映当前二次水动力模型下、所列移动时间块口径中的训练样本与残差补全不确定性。",
        "- 未纳入流量/水位仪器误差、模型族选择、长时间相关结构及未观测洪峰的系统偏差。",
        "- 最近邻映射保持了源块内残差的时间次序和经验取值，但稀疏实测无法恢复块内未观测的连续残差轨迹。",
        "- 因 87% 左右含沙量由模型补全，年度排沙量应连同区间、场景敏感性和上述限制一起报告。",
        "- 相关系数仅为描述性统计；没有显著性或因果识别证据时，不使用“显著影响”措辞。",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    args = parse_args()
    input_path = resolve_cli_path(args.input)
    run_dir = resolve_cli_path(args.run_dir)
    validate_output_location(run_dir)
    block_hours_values = validate_block_hours(
        args.canonical_block_hours, args.block_sensitivity_hours
    )
    if not input_path.is_file():
        raise FileNotFoundError(input_path)

    results_data_dir = run_dir / "results" / "data"
    results_tables_dir = run_dir / "results" / "tables"
    validation_dir = run_dir / "validation"
    logs_dir = run_dir / "logs"
    for directory in (results_data_dir, results_tables_dir, validation_dir, logs_dir):
        directory.mkdir(parents=True, exist_ok=True)

    df = load_hydro_data(input_path)
    observed = df.dropna(subset=["flow_m3s", "water_level_m", "sediment_obs_kgm3"]).copy()
    observed = observed[observed["sediment_obs_kgm3"] > 0].copy().sort_values("datetime")

    validation = year_holdout_cv(observed, MODEL_KINDS)
    selected_model_name = str(
        validation[validation["heldout_year"].eq("mean")]
        .sort_values("rmse_log")
        .iloc[0]["model"]
    )
    selected_kind = MODEL_KINDS[selected_model_name]
    x_train, names = build_features(observed, selected_kind)
    final_model = fit_linear_model(
        selected_model_name,
        x_train,
        np.log(observed["sediment_obs_kgm3"].to_numpy(dtype=float)),
        names,
    )
    completed = build_completed_series(df, final_model, selected_kind, model_bounds(observed))
    annual = summarize_annual(completed)
    monthly = summarize_monthly(completed)
    relationships = summarize_relationships(observed, final_model, selected_kind)
    holdout = annual_holdout_flux_validation(observed, selected_model_name, selected_kind)
    sensitivity, sensitivity_summary = build_sensitivity(
        df, observed, annual, selected_model_name, selected_kind, final_model
    )
    bootstrap_draw_parts = []
    uncertainty_parts = []
    for block_hours in block_hours_values:
        draws_part, uncertainty_part = bootstrap_annual_uncertainty(
            df,
            observed,
            annual,
            selected_model_name,
            selected_kind,
            args.bootstrap_replicates,
            args.seed,
            block_hours,
        )
        draws_part["is_canonical"] = block_hours == args.canonical_block_hours
        uncertainty_part["is_canonical"] = block_hours == args.canonical_block_hours
        bootstrap_draw_parts.append(draws_part)
        uncertainty_parts.append(uncertainty_part)
    bootstrap_block_draws = pd.concat(bootstrap_draw_parts, ignore_index=True)
    block_sensitivity = pd.concat(uncertainty_parts, ignore_index=True)
    bootstrap_draws = bootstrap_block_draws[
        bootstrap_block_draws["is_canonical"]
    ].reset_index(drop=True)
    uncertainty = block_sensitivity[block_sensitivity["is_canonical"]].reset_index(drop=True)
    block_diagnostics = build_time_block_diagnostics(
        observed, block_hours_values, args.canonical_block_hours
    )
    quality_checks = build_quality_checks(
        completed,
        annual,
        uncertainty,
        block_sensitivity,
        block_diagnostics,
        selected_model_name,
        args.canonical_block_hours,
    )

    completed.to_csv(
        results_data_dir / "cleaned_hydro_timeseries.csv", index=False, encoding="utf-8-sig"
    )
    annual.to_csv(results_tables_dir / "annual_flux_estimates.csv", index=False, encoding="utf-8-sig")
    monthly.to_csv(
        results_tables_dir / "monthly_relationship_summary.csv", index=False, encoding="utf-8-sig"
    )
    validation.to_csv(validation_dir / "model_validation.csv", index=False, encoding="utf-8-sig")
    relationships.to_csv(
        validation_dir / "relationship_diagnostics.csv", index=False, encoding="utf-8-sig"
    )
    holdout.to_csv(
        validation_dir / "annual_flux_holdout_validation.csv", index=False, encoding="utf-8-sig"
    )
    sensitivity.to_csv(
        validation_dir / "annual_sensitivity_scenarios.csv", index=False, encoding="utf-8-sig"
    )
    sensitivity_summary.to_csv(
        validation_dir / "annual_sensitivity_summary.csv", index=False, encoding="utf-8-sig"
    )
    bootstrap_draws.to_csv(
        validation_dir / "bootstrap_annual_draws.csv", index=False, encoding="utf-8-sig"
    )
    bootstrap_block_draws.to_csv(
        validation_dir / "bootstrap_block_sensitivity_draws.csv",
        index=False,
        encoding="utf-8-sig",
    )
    uncertainty.to_csv(
        validation_dir / "annual_sediment_uncertainty.csv", index=False, encoding="utf-8-sig"
    )
    block_sensitivity.to_csv(
        validation_dir / "bootstrap_block_length_sensitivity.csv",
        index=False,
        encoding="utf-8-sig",
    )
    block_diagnostics.to_csv(
        validation_dir / "bootstrap_time_block_diagnostics.csv",
        index=False,
        encoding="utf-8-sig",
    )
    quality_checks.to_csv(
        validation_dir / "quality_checks.csv", index=False, encoding="utf-8-sig"
    )
    write_analysis(
        run_dir / "results" / "analysis.md",
        input_path,
        completed,
        validation,
        annual,
        holdout,
        sensitivity_summary,
        uncertainty,
        block_sensitivity,
        selected_model_name,
        names,
        args.canonical_block_hours,
    )
    write_validation_summary(
        validation_dir / "validation_summary.md",
        holdout,
        sensitivity_summary,
        uncertainty,
        block_sensitivity,
        block_diagnostics,
        args.bootstrap_replicates,
        args.seed,
        args.canonical_block_hours,
        block_hours_values,
    )

    log_lines = [
        f"Input: {input_path}",
        f"Run directory: {run_dir}",
        f"Python: {platform.python_version()}",
        f"NumPy: {np.__version__}",
        f"pandas: {pd.__version__}",
        f"openpyxl: {openpyxl.__version__}",
        f"Cleaned rows: {len(completed)}",
        f"Observed sediment rows: {len(observed)}",
        f"Modeled sediment rows: {completed['sediment_filled_flag'].sum()}",
        f"Selected machine candidate: {selected_model_name}",
        f"Bootstrap replicates: {args.bootstrap_replicates}",
        f"Random seed: {args.seed}",
        f"Canonical moving-time block hours: {args.canonical_block_hours}",
        f"Block sensitivity hours: {','.join(str(value) for value in block_hours_values)}",
        "Pairs resampling: within-year non-circular moving real-time blocks",
        "Residual resampling: within-year moving real-time blocks mapped by relative time",
        f"Quality checks passed: {int(quality_checks['passed'].sum())}/{len(quality_checks)}",
    ]
    (logs_dir / "run_summary.txt").write_text("\n".join(log_lines) + "\n", encoding="utf-8")
    print("\n".join(log_lines))
    print("Annual estimates:")
    print(
        annual[["year", "water_volume_1e8_m3", "sediment_mass_1e4_t"]].to_string(index=False)
    )
    if not quality_checks["passed"].all():
        failed = quality_checks.loc[~quality_checks["passed"], "check"].tolist()
        raise RuntimeError(f"Quality checks failed: {failed}")


if __name__ == "__main__":
    main()
