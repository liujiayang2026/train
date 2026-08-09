from __future__ import annotations

import html
import math
import re
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd


SEED = 20220803
N_ASSOC_PERM = 50_000
N_COMPONENT_PERM = 20_000
N_FACTOR_BOOTSTRAP = 5_000
N_COEFFICIENT_BOOTSTRAP = 2_000
PSEUDOCOUNT = 0.10
EFFECT_CAP = 2.50

COMPONENTS = [
    "SiO2", "Na2O", "K2O", "CaO", "MgO", "Al2O3", "Fe2O3",
    "CuO", "PbO", "BaO", "P2O5", "SrO", "SnO2", "SO2",
]
COMPONENT_LABELS = {
    "SiO2": "SiO2", "Na2O": "Na2O", "K2O": "K2O", "CaO": "CaO",
    "MgO": "MgO", "Al2O3": "Al2O3", "Fe2O3": "Fe2O3", "CuO": "CuO",
    "PbO": "PbO", "BaO": "BaO", "P2O5": "P2O5", "SrO": "SrO",
    "SnO2": "SnO2", "SO2": "SO2",
}


def find_source_workbook(project_root: Path) -> Path:
    candidates = [
        path for path in project_root.rglob("*.xlsx")
        if not path.name.startswith("~$") and "question-1" not in path.parts
    ]
    if len(candidates) != 1:
        raise RuntimeError(f"Expected one source workbook, found: {candidates}")
    return candidates[0]


def read_workbook(path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    artifact_rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    sample_rows = list(workbook.worksheets[1].iter_rows(values_only=True))

    artifacts = pd.DataFrame(artifact_rows[1:], columns=[
        "artifact_id", "pattern", "glass_type", "color", "surface_weathering"
    ])
    artifacts["artifact_id"] = artifacts["artifact_id"].astype(str).str.zfill(2)
    artifacts["color"] = artifacts["color"].fillna("未知")

    samples = pd.DataFrame(sample_rows[1:], columns=["sample_id", *COMPONENTS])
    samples["sample_id"] = samples["sample_id"].astype(str)
    samples["artifact_id"] = samples["sample_id"].str.extract(r"^(\d+)", expand=False).str.zfill(2)
    samples[COMPONENTS] = samples[COMPONENTS].apply(pd.to_numeric, errors="coerce").fillna(0.0)
    samples["component_sum"] = samples[COMPONENTS].sum(axis=1)
    samples["is_valid"] = samples["component_sum"].between(85.0, 105.0)
    samples = samples.merge(artifacts, on="artifact_id", how="left", validate="many_to_one")
    samples["point_weathering"] = samples["surface_weathering"]
    samples.loc[samples["sample_id"].str.contains("未风化", regex=False), "point_weathering"] = "无风化"
    samples.loc[samples["sample_id"].str.contains("严重风化", regex=False), "point_weathering"] = "风化"
    return artifacts, samples


def chi_square_stat(table: np.ndarray) -> float:
    table = np.asarray(table, dtype=float)
    expected = np.outer(table.sum(axis=1), table.sum(axis=0)) / table.sum()
    return float(np.sum((table - expected) ** 2 / np.where(expected > 0, expected, 1)))


def association_test(feature: pd.Series, weather: pd.Series, rng: np.random.Generator) -> dict:
    x = feature.astype(str).to_numpy()
    y = (weather == "风化").astype(int).to_numpy()
    categories, x_codes = np.unique(x, return_inverse=True)
    table = np.zeros((len(categories), 2), dtype=int)
    np.add.at(table, (x_codes, y), 1)
    observed = chi_square_stat(table)
    exceed = 0
    for _ in range(N_ASSOC_PERM):
        yp = rng.permutation(y)
        perm_table = np.zeros_like(table)
        np.add.at(perm_table, (x_codes, yp), 1)
        exceed += chi_square_stat(perm_table) >= observed - 1e-12
    p_value = (exceed + 1) / (N_ASSOC_PERM + 1)
    denom = max(1, min(table.shape[0] - 1, table.shape[1] - 1))
    cramers_v = math.sqrt(observed / (table.sum() * denom))
    return {
        "chi_square": observed,
        "degrees_of_freedom": int((table.shape[0] - 1) * (table.shape[1] - 1)),
        "permutation_p": p_value,
        "cramers_v": cramers_v,
        "levels": len(categories),
    }


def sigmoid(values: np.ndarray) -> np.ndarray:
    return 1.0 / (1.0 + np.exp(-np.clip(values, -35.0, 35.0)))


def logistic_log_likelihood(y: np.ndarray, probabilities: np.ndarray) -> float:
    probabilities = np.clip(probabilities, 1e-12, 1.0 - 1e-12)
    return float(np.sum(y * np.log(probabilities) + (1.0 - y) * np.log(1.0 - probabilities)))


def firth_penalized_log_likelihood(X: np.ndarray, y: np.ndarray, beta: np.ndarray) -> float:
    probabilities = sigmoid(X @ beta)
    weights = np.maximum(probabilities * (1.0 - probabilities), 1e-10)
    information = X.T @ (weights[:, None] * X)
    sign, log_determinant = np.linalg.slogdet(information)
    if sign <= 0:
        return -np.inf
    return logistic_log_likelihood(y, probabilities) + 0.5 * log_determinant


def fit_firth_logistic(
    X: np.ndarray,
    y: np.ndarray,
    initial: np.ndarray | None = None,
    max_iterations: int = 160,
    tolerance: float = 1e-6,
) -> dict:
    X = np.asarray(X, dtype=float)
    y = np.asarray(y, dtype=float)
    beta = np.zeros(X.shape[1], dtype=float) if initial is None else np.asarray(initial, dtype=float).copy()
    converged = False
    penalized_ll = firth_penalized_log_likelihood(X, y, beta)
    for iteration in range(1, max_iterations + 1):
        probabilities = sigmoid(X @ beta)
        weights = np.maximum(probabilities * (1.0 - probabilities), 1e-10)
        information = X.T @ (weights[:, None] * X)
        inverse_information = np.linalg.pinv(information, rcond=1e-12)
        leverage = weights * np.sum((X @ inverse_information) * X, axis=1)
        adjusted_score = X.T @ (y - probabilities + leverage * (0.5 - probabilities))
        step = inverse_information @ adjusted_score
        step_scale = 1.0
        accepted = False
        while step_scale >= 2.0 ** -20:
            candidate = beta + step_scale * step
            candidate_ll = firth_penalized_log_likelihood(X, y, candidate)
            if candidate_ll >= penalized_ll - 1e-10:
                accepted = True
                break
            step_scale *= 0.5
        if not accepted:
            break
        beta = candidate
        penalized_ll = candidate_ll
        if np.max(np.abs(step_scale * step)) < tolerance:
            converged = True
            break
    probabilities = sigmoid(X @ beta)
    weights = np.maximum(probabilities * (1.0 - probabilities), 1e-10)
    information = X.T @ (weights[:, None] * X)
    inverse_information = np.linalg.pinv(information, rcond=1e-12)
    leverage = weights * np.sum((X @ inverse_information) * X, axis=1)
    final_score = X.T @ (y - probabilities + leverage * (0.5 - probabilities))
    converged = converged or np.max(np.abs(final_score)) < 1e-4
    return {
        "beta": beta,
        "probabilities": probabilities,
        "log_likelihood": logistic_log_likelihood(y, probabilities),
        "penalized_log_likelihood": penalized_ll,
        "covariance": inverse_information,
        "maximum_adjusted_score": float(np.max(np.abs(final_score))),
        "converged": converged,
        "iterations": iteration,
    }


def build_weathering_design(artifacts: pd.DataFrame) -> tuple[np.ndarray, np.ndarray, list[str], dict]:
    references = {"glass_type": "高钾", "pattern": "A", "color": "浅蓝"}
    factors = ("glass_type", "pattern", "color")
    columns = [np.ones(len(artifacts), dtype=float)]
    names = ["intercept"]
    factor_columns = {}
    level_metadata = {}
    for factor in factors:
        levels = sorted(artifacts[factor].astype(str).unique())
        reference = references[factor]
        if reference not in levels:
            raise RuntimeError(f"Reference level {reference} is missing for {factor}")
        factor_columns[factor] = []
        level_metadata[factor] = {"reference": reference, "levels": levels, "columns": {}}
        for level in levels:
            if level == reference:
                continue
            column_index = len(columns)
            columns.append(artifacts[factor].astype(str).eq(level).to_numpy(float))
            names.append(f"{factor}={level}")
            factor_columns[factor].append(column_index)
            level_metadata[factor]["columns"][level] = column_index
    X = np.column_stack(columns)
    y = artifacts["surface_weathering"].eq("风化").to_numpy(float)
    metadata = {
        "factor_columns": factor_columns,
        "levels": level_metadata,
        "references": references,
    }
    return X, y, names, metadata


def firth_global_factor_tests(
    X: np.ndarray,
    y: np.ndarray,
    metadata: dict,
    rng: np.random.Generator,
) -> tuple[pd.DataFrame, dict]:
    full_fit = fit_firth_logistic(X, y)
    if not full_fit["converged"]:
        raise RuntimeError("Observed full Firth model did not converge")
    rows = []
    all_columns = np.arange(X.shape[1])
    for factor, removed_columns in metadata["factor_columns"].items():
        keep_columns = np.array([column for column in all_columns if column not in removed_columns])
        reduced_X = X[:, keep_columns]
        reduced_fit = fit_firth_logistic(reduced_X, y)
        if not reduced_fit["converged"]:
            raise RuntimeError(f"Observed reduced Firth model did not converge for {factor}")
        observed_statistic = max(
            0.0,
            2.0 * (full_fit["log_likelihood"] - reduced_fit["log_likelihood"]),
        )
        exceed = 0
        successful = 0
        for _ in range(N_FACTOR_BOOTSTRAP):
            simulated_y = rng.binomial(1, reduced_fit["probabilities"]).astype(float)
            simulated_full = fit_firth_logistic(X, simulated_y, initial=full_fit["beta"])
            simulated_reduced = fit_firth_logistic(
                reduced_X, simulated_y, initial=reduced_fit["beta"]
            )
            if not (simulated_full["converged"] and simulated_reduced["converged"]):
                continue
            simulated_statistic = max(
                0.0,
                2.0 * (
                    simulated_full["log_likelihood"] - simulated_reduced["log_likelihood"]
                ),
            )
            exceed += simulated_statistic >= observed_statistic - 1e-12
            successful += 1
        if successful < 0.95 * N_FACTOR_BOOTSTRAP:
            raise RuntimeError(f"Too many failed bootstrap fits for {factor}: {successful}")
        rows.append({
            "factor": factor,
            "observed_deviance": observed_statistic,
            "parameter_bootstrap_p": (exceed + 1) / (successful + 1),
            "requested_simulations": N_FACTOR_BOOTSTRAP,
            "successful_simulations": successful,
            "reduced_model_iterations": reduced_fit["iterations"],
        })
    return pd.DataFrame(rows), full_fit


def bootstrap_firth_coefficients(
    X: np.ndarray,
    full_fit: dict,
    rng: np.random.Generator,
) -> tuple[np.ndarray, pd.DataFrame]:
    estimates = []
    failed = 0
    for _ in range(N_COEFFICIENT_BOOTSTRAP):
        simulated_y = rng.binomial(1, full_fit["probabilities"]).astype(float)
        simulated_fit = fit_firth_logistic(X, simulated_y, initial=full_fit["beta"])
        if simulated_fit["converged"]:
            estimates.append(simulated_fit["beta"])
        else:
            failed += 1
    if len(estimates) < 0.95 * N_COEFFICIENT_BOOTSTRAP:
        raise RuntimeError(f"Too many failed coefficient bootstrap fits: {failed}")
    diagnostics = pd.DataFrame([{
        "requested_simulations": N_COEFFICIENT_BOOTSTRAP,
        "successful_simulations": len(estimates),
        "failed_simulations": failed,
        "full_model_iterations": full_fit["iterations"],
        "full_model_converged": full_fit["converged"],
    }])
    return np.vstack(estimates), diagnostics


def firth_result_tables(
    X: np.ndarray,
    names: list[str],
    metadata: dict,
    full_fit: dict,
    bootstrap_betas: np.ndarray,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    coefficient_rows = []
    for factor, factor_info in metadata["levels"].items():
        reference = factor_info["reference"]
        for level, column_index in factor_info["columns"].items():
            estimate = full_fit["beta"][column_index]
            bootstrap_values = bootstrap_betas[:, column_index]
            coefficient_rows.append({
                "factor": factor,
                "level": level,
                "reference": reference,
                "term": names[column_index],
                "coefficient": estimate,
                "adjusted_odds_ratio": math.exp(estimate),
                "or_ci_low": math.exp(np.quantile(bootstrap_values, 0.025)),
                "or_ci_high": math.exp(np.quantile(bootstrap_values, 0.975)),
                "positive_coefficient_probability": np.mean(bootstrap_values > 0.0),
            })
    coefficient_table = pd.DataFrame(coefficient_rows)

    probability_rows = []
    beta = full_fit["beta"]
    for factor, factor_info in metadata["levels"].items():
        factor_columns = metadata["factor_columns"][factor]
        reference = factor_info["reference"]
        reference_X = X.copy()
        reference_X[:, factor_columns] = 0.0
        reference_probability = sigmoid(reference_X @ beta).mean()
        reference_bootstrap = sigmoid(reference_X @ bootstrap_betas.T).mean(axis=0)
        for level in factor_info["levels"]:
            counterfactual_X = reference_X.copy()
            if level != reference:
                counterfactual_X[:, factor_info["columns"][level]] = 1.0
            adjusted_probability = sigmoid(counterfactual_X @ beta).mean()
            bootstrap_probability = sigmoid(counterfactual_X @ bootstrap_betas.T).mean(axis=0)
            bootstrap_effect = bootstrap_probability - reference_bootstrap
            probability_rows.append({
                "factor": factor,
                "level": level,
                "reference": reference,
                "adjusted_probability": adjusted_probability,
                "probability_ci_low": np.quantile(bootstrap_probability, 0.025),
                "probability_ci_high": np.quantile(bootstrap_probability, 0.975),
                "average_marginal_effect_vs_reference": adjusted_probability - reference_probability,
                "ame_ci_low": np.quantile(bootstrap_effect, 0.025),
                "ame_ci_high": np.quantile(bootstrap_effect, 0.975),
            })
    return coefficient_table, pd.DataFrame(probability_rows)


def firth_specification_sensitivity(
    X: np.ndarray,
    y: np.ndarray,
    names: list[str],
    metadata: dict,
) -> pd.DataFrame:
    specifications = [
        ("type-only", ["glass_type"]),
        ("type-pattern", ["glass_type", "pattern"]),
        ("type-color", ["glass_type", "color"]),
        ("full", ["glass_type", "pattern", "color"]),
    ]
    rows = []
    for specification, factors in specifications:
        selected = [0] + [
            column
            for factor in factors
            for column in metadata["factor_columns"][factor]
        ]
        fit = fit_firth_logistic(X[:, selected], y)
        local_index = {global_index: index for index, global_index in enumerate(selected)}
        row = {
            "specification": specification,
            "factors": "+".join(factors),
            "number_of_parameters": len(selected),
            "converged": fit["converged"],
            "log_likelihood": fit["log_likelihood"],
        }
        for factor in ("glass_type", "pattern"):
            for global_index in metadata["factor_columns"][factor]:
                term = names[global_index]
                output_name = "or_" + term.replace("=", "_")
                row[output_name] = (
                    math.exp(fit["beta"][local_index[global_index]])
                    if global_index in local_index else np.nan
                )
        rows.append(row)
    return pd.DataFrame(rows)


def benjamini_hochberg(values: np.ndarray) -> np.ndarray:
    values = np.asarray(values, dtype=float)
    order = np.argsort(values)
    adjusted = np.empty_like(values)
    ranked = values[order] * len(values) / np.arange(1, len(values) + 1)
    ranked = np.minimum.accumulate(ranked[::-1])[::-1]
    adjusted[order] = np.minimum(ranked, 1.0)
    return adjusted


def component_statistics(units: pd.DataFrame, rng: np.random.Generator) -> pd.DataFrame:
    rows = []
    for glass_type in sorted(units["glass_type"].unique()):
        subset = units.loc[units["glass_type"] == glass_type].reset_index(drop=True)
        weather_mask = subset["point_weathering"].eq("风化").to_numpy()
        current = subset[COMPONENTS].to_numpy(float)
        observed = np.abs(current[weather_mask].mean(axis=0) - current[~weather_mask].mean(axis=0))
        exceed = np.zeros(len(COMPONENTS), dtype=int)
        for _ in range(N_COMPONENT_PERM):
            perm_mask = rng.permutation(weather_mask)
            perm_diff = np.abs(current[perm_mask].mean(axis=0) - current[~perm_mask].mean(axis=0))
            exceed += perm_diff >= observed - 1e-12
        p_values = (exceed + 1) / (N_COMPONENT_PERM + 1)
        q_values = benjamini_hochberg(p_values)
        for index, component in enumerate(COMPONENTS):
            weather_values = current[weather_mask, index]
            clear_values = current[~weather_mask, index]
            rows.append({
                "glass_type": glass_type,
                "component": component,
                "n_weathered": len(weather_values),
                "n_unweathered": len(clear_values),
                "weathered_mean": weather_values.mean(),
                "unweathered_mean": clear_values.mean(),
                "mean_difference": weather_values.mean() - clear_values.mean(),
                "weathered_median": np.median(weather_values),
                "unweathered_median": np.median(clear_values),
                "permutation_p": p_values[index],
                "bh_q": q_values[index],
            })
    return pd.DataFrame(rows)


def close_composition(values: np.ndarray) -> np.ndarray:
    values = np.asarray(values, dtype=float)
    return values / values.sum(axis=1, keepdims=True) * 100.0


def clr_transform(values: np.ndarray, pseudocount: float) -> np.ndarray:
    closed = close_composition(values + pseudocount)
    logged = np.log(closed)
    return logged - logged.mean(axis=1, keepdims=True)


def inverse_clr(values: np.ndarray) -> np.ndarray:
    shifted = values - values.max(axis=1, keepdims=True)
    exponentiated = np.exp(shifted)
    return close_composition(exponentiated)


def estimate_weathering_effect(units: pd.DataFrame, pseudocount: float) -> dict[str, np.ndarray]:
    effects = {}
    for glass_type, subset in units.groupby("glass_type"):
        transformed = clr_transform(subset[COMPONENTS].to_numpy(float), pseudocount)
        weather = subset["point_weathering"].eq("风化").to_numpy()
        effect = np.median(transformed[weather], axis=0) - np.median(transformed[~weather], axis=0)
        effects[glass_type] = np.clip(effect, -EFFECT_CAP, EFFECT_CAP)
    return effects


def predict_preweathering(
    valid_samples: pd.DataFrame,
    units: pd.DataFrame,
    pseudocount: float,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    effects = estimate_weathering_effect(units, pseudocount)
    weathered = valid_samples.loc[valid_samples["point_weathering"] == "风化"].copy()
    predicted_blocks = []
    effect_rows = []
    for glass_type, subset in weathered.groupby("glass_type", sort=True):
        transformed = clr_transform(subset[COMPONENTS].to_numpy(float), pseudocount)
        predicted = inverse_clr(transformed - effects[glass_type])
        block = subset[["sample_id", "artifact_id", "glass_type", "pattern", "color", "component_sum"]].copy()
        for index, component in enumerate(COMPONENTS):
            block[f"current_{component}"] = subset[component].to_numpy(float)
            block[f"predicted_{component}"] = predicted[:, index]
            effect_rows.append({
                "glass_type": glass_type,
                "component": component,
                "clr_weathering_effect": effects[glass_type][index],
                "multiplicative_direction": math.exp(effects[glass_type][index]),
                "effect_was_capped": abs(effects[glass_type][index]) >= EFFECT_CAP - 1e-12,
            })
        block["predicted_sum"] = predicted.sum(axis=1)
        predicted_blocks.append(block)
    return pd.concat(predicted_blocks, ignore_index=True), pd.DataFrame(effect_rows).drop_duplicates()


def sensitivity_analysis(valid_samples: pd.DataFrame, units: pd.DataFrame) -> pd.DataFrame:
    predictions = {}
    for pseudocount in (0.05, 0.10, 0.50):
        predicted, _ = predict_preweathering(valid_samples, units, pseudocount)
        predicted = predicted.set_index("sample_id")
        predictions[pseudocount] = predicted[[f"predicted_{c}" for c in COMPONENTS]]
    base = predictions[0.10]
    rows = []
    for pseudocount, values in predictions.items():
        difference = np.abs(values.to_numpy() - base.to_numpy())
        rows.append({
            "pseudocount": pseudocount,
            "mean_absolute_change_vs_0_10": difference.mean(),
            "maximum_absolute_change_vs_0_10": difference.max(),
            "p95_absolute_change_vs_0_10": np.quantile(difference, 0.95),
        })
    return pd.DataFrame(rows)


def validate_paired_artifacts(units: pd.DataFrame, predictions: pd.DataFrame) -> pd.DataFrame:
    state_counts = units.groupby("artifact_id")["point_weathering"].nunique()
    paired_ids = state_counts[state_counts > 1].index
    rows = []
    for artifact_id in paired_ids:
        observed = units.loc[
            (units["artifact_id"] == artifact_id) & (units["point_weathering"] == "无风化"),
            COMPONENTS,
        ].mean().to_numpy(float)
        predicted_rows = predictions.loc[predictions["artifact_id"] == artifact_id]
        current = predicted_rows[[f"current_{c}" for c in COMPONENTS]].mean().to_numpy(float)
        predicted = predicted_rows[[f"predicted_{c}" for c in COMPONENTS]].mean().to_numpy(float)
        rows.append({
            "artifact_id": artifact_id,
            "glass_type": predicted_rows["glass_type"].iloc[0],
            "current_mae": np.abs(current - observed).mean(),
            "predicted_mae": np.abs(predicted - observed).mean(),
            "mae_improvement": np.abs(current - observed).mean() - np.abs(predicted - observed).mean(),
            "current_rmse": np.sqrt(np.mean((current - observed) ** 2)),
            "predicted_rmse": np.sqrt(np.mean((predicted - observed) ** 2)),
        })
    return pd.DataFrame(rows)


def svg_start(width: int, height: int, title: str) -> list[str]:
    return [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<rect width="100%" height="100%" fill="#ffffff"/>',
        '<style>text{font-family:"Microsoft YaHei",Arial,sans-serif;fill:#17212b;letter-spacing:0}.title{font-size:20px;font-weight:700}.axis{font-size:12px}.label{font-size:13px}.small{font-size:11px}</style>',
        f'<text x="{width/2}" y="30" text-anchor="middle" class="title">{html.escape(title)}</text>',
    ]


def write_association_svg(artifacts: pd.DataFrame, path: Path) -> None:
    panels = [("glass_type", "玻璃类型"), ("pattern", "纹饰"), ("color", "颜色")]
    width, panel_width, height = 1260, 400, 500
    parts = svg_start(width, height, "不同文物特征的表面风化率")
    for panel_index, (column, panel_title) in enumerate(panels):
        x0 = 30 + panel_index * panel_width
        rates = artifacts.groupby(column)["surface_weathering"].apply(lambda x: (x == "风化").mean()).sort_values(ascending=False)
        counts = artifacts.groupby(column).size().reindex(rates.index)
        parts.append(f'<text x="{x0 + 180}" y="62" text-anchor="middle" class="label" font-weight="700">{panel_title}</text>')
        bar_area = 310
        row_height = min(35, 370 / max(1, len(rates)))
        for i, (label, rate) in enumerate(rates.items()):
            y = 82 + i * row_height
            bar_width = bar_area * rate
            parts.append(f'<text x="{x0}" y="{y + 15}" class="axis">{html.escape(str(label))}</text>')
            parts.append(f'<rect x="{x0 + 62}" y="{y}" width="{bar_area}" height="18" fill="#e8edf1"/>')
            parts.append(f'<rect x="{x0 + 62}" y="{y}" width="{bar_width:.1f}" height="18" fill="#2f7d6d"/>')
            parts.append(f'<text x="{x0 + 68 + bar_width:.1f}" y="{y + 14}" class="small">{rate:.1%} (n={counts[label]})</text>')
    parts.append('</svg>')
    path.write_text("\n".join(parts), encoding="utf-8")


def color_for_difference(value: float, bound: float) -> str:
    ratio = min(abs(value) / bound, 1.0)
    if value >= 0:
        start, end = (247, 248, 248), (188, 67, 52)
    else:
        start, end = (247, 248, 248), (42, 110, 145)
    rgb = tuple(round(start[i] + ratio * (end[i] - start[i])) for i in range(3))
    return f"rgb{rgb}"


def write_component_heatmap(stats: pd.DataFrame, path: Path) -> None:
    pivot = stats.pivot(index="glass_type", columns="component", values="mean_difference").reindex(columns=COMPONENTS)
    width, height = 1100, 260
    parts = svg_start(width, height, "风化组与无风化组的平均成分差（百分点）")
    x0, y0, cell_w, cell_h = 130, 92, 64, 48
    bound = max(1.0, float(np.quantile(np.abs(pivot.to_numpy()), 0.90)))
    for j, component in enumerate(COMPONENTS):
        parts.append(f'<text x="{x0 + j*cell_w + cell_w/2}" y="78" text-anchor="middle" class="axis">{component}</text>')
    for i, (glass_type, row) in enumerate(pivot.iterrows()):
        parts.append(f'<text x="{x0 - 12}" y="{y0+i*cell_h+29}" text-anchor="end" class="label">{html.escape(str(glass_type))}</text>')
        for j, value in enumerate(row):
            color = color_for_difference(float(value), bound)
            x, y = x0 + j * cell_w, y0 + i * cell_h
            parts.append(f'<rect x="{x}" y="{y}" width="{cell_w-2}" height="{cell_h-2}" fill="{color}"/>')
            text_color = "#ffffff" if abs(value) / bound > 0.55 else "#17212b"
            parts.append(f'<text x="{x+cell_w/2-1}" y="{y+29}" text-anchor="middle" class="small" fill="{text_color}" style="fill:{text_color}">{value:+.2f}</text>')
    parts.append('<text x="550" y="230" text-anchor="middle" class="small">红色表示风化后升高，蓝色表示风化后降低；颜色按第90百分位缩放</text>')
    parts.append('</svg>')
    path.write_text("\n".join(parts), encoding="utf-8")


def write_prediction_svg(predictions: pd.DataFrame, path: Path) -> None:
    width, height = 1120, 520
    parts = svg_start(width, height, "风化点当前成分与预测风化前成分均值")
    types = sorted(predictions["glass_type"].unique())
    for panel_index, glass_type in enumerate(types):
        subset = predictions[predictions["glass_type"] == glass_type]
        current = np.array([subset[f"current_{c}"].mean() for c in COMPONENTS])
        predicted = np.array([subset[f"predicted_{c}"].mean() for c in COMPONENTS])
        top = np.argsort(current + predicted)[-8:][::-1]
        x0 = 60 + panel_index * 540
        baseline, max_height = 445, 310
        maximum = max(float(current[top].max()), float(predicted[top].max()), 1.0)
        parts.append(f'<text x="{x0+230}" y="64" text-anchor="middle" class="label" font-weight="700">{html.escape(glass_type)}</text>')
        for slot, index in enumerate(top):
            x = x0 + slot * 57
            h1 = max_height * current[index] / maximum
            h2 = max_height * predicted[index] / maximum
            parts.append(f'<rect x="{x}" y="{baseline-h1:.1f}" width="20" height="{h1:.1f}" fill="#b94b3e"/>')
            parts.append(f'<rect x="{x+21}" y="{baseline-h2:.1f}" width="20" height="{h2:.1f}" fill="#2f7d6d"/>')
            parts.append(f'<text x="{x+20}" y="{baseline+18}" text-anchor="middle" class="small">{COMPONENTS[index]}</text>')
        parts.append(f'<line x1="{x0-8}" y1="{baseline}" x2="{x0+460}" y2="{baseline}" stroke="#65727e"/>')
    parts.append('<rect x="430" y="480" width="15" height="12" fill="#b94b3e"/><text x="452" y="491" class="axis">当前风化点</text>')
    parts.append('<rect x="550" y="480" width="15" height="12" fill="#2f7d6d"/><text x="572" y="491" class="axis">预测风化前</text>')
    parts.append('</svg>')
    path.write_text("\n".join(parts), encoding="utf-8")


def write_adjusted_probability_svg(probabilities: pd.DataFrame, path: Path) -> None:
    factor_labels = {"glass_type": "玻璃类型", "pattern": "纹饰", "color": "颜色"}
    factor_colors = {"glass_type": "#2f7d6d", "pattern": "#b94b3e", "color": "#3c6e9b"}
    plot = probabilities.copy()
    width = 1050
    row_height = 34
    height = 100 + row_height * len(plot) + 45
    x0, plot_width = 250, 700
    parts = svg_start(width, height, "Firth模型调整后的风化概率及95% Bootstrap区间")
    for tick in np.linspace(0.0, 1.0, 6):
        x = x0 + tick * plot_width
        parts.append(f'<line x1="{x:.1f}" y1="64" x2="{x:.1f}" y2="{height-45}" stroke="#e1e6ea"/>')
        parts.append(f'<text x="{x:.1f}" y="{height-22}" text-anchor="middle" class="axis">{tick:.0%}</text>')
    for row_index, row in enumerate(plot.itertuples()):
        y = 78 + row_index * row_height
        label = f"{factor_labels[row.factor]}：{row.level}"
        low = x0 + row.probability_ci_low * plot_width
        high = x0 + row.probability_ci_high * plot_width
        point = x0 + row.adjusted_probability * plot_width
        color = factor_colors[row.factor]
        parts.append(f'<text x="{x0-16}" y="{y+5}" text-anchor="end" class="label">{html.escape(label)}</text>')
        parts.append(f'<line x1="{low:.1f}" y1="{y}" x2="{high:.1f}" y2="{y}" stroke="{color}" stroke-width="3"/>')
        parts.append(f'<line x1="{low:.1f}" y1="{y-5}" x2="{low:.1f}" y2="{y+5}" stroke="{color}"/>')
        parts.append(f'<line x1="{high:.1f}" y1="{y-5}" x2="{high:.1f}" y2="{y+5}" stroke="{color}"/>')
        parts.append(f'<circle cx="{point:.1f}" cy="{y}" r="5" fill="{color}"/>')
        parts.append(f'<text x="{min(point+10, width-70):.1f}" y="{y+5}" class="small">{row.adjusted_probability:.1%}</text>')
    parts.append('</svg>')
    path.write_text("\n".join(parts), encoding="utf-8")


def build_report(
    artifacts: pd.DataFrame,
    samples: pd.DataFrame,
    association: pd.DataFrame,
    firth_global: pd.DataFrame,
    firth_coefficients: pd.DataFrame,
    adjusted_probabilities: pd.DataFrame,
    specification_sensitivity: pd.DataFrame,
    stats: pd.DataFrame,
    predictions: pd.DataFrame,
    effects: pd.DataFrame,
    sensitivity: pd.DataFrame,
    paired_validation: pd.DataFrame,
) -> str:
    invalid = samples.loc[~samples["is_valid"], ["sample_id", "component_sum"]]
    significant = stats.loc[stats["bh_q"] < 0.05].sort_values(["glass_type", "bh_q"])
    lines = [
        "# Question 1 preliminary analysis",
        "",
        "## Data checks",
        "",
        f"- Artifact records: {len(artifacts)}.",
        f"- Sampling-point records: {len(samples)}; valid records: {int(samples['is_valid'].sum())}.",
        "- Invalid points: " + ", ".join(f"{r.sample_id} (sum={r.component_sum:.2f})" for r in invalid.itertuples()) + ".",
        "- Missing component cells were treated as zero. Point names containing `未风化` override artifact-level weathering as unweathered.",
        "",
        "## Association with surface weathering",
        "",
    ]
    feature_names = {"glass_type": "玻璃类型", "pattern": "纹饰", "color": "颜色"}
    lines.append("### Unadjusted permutation tests")
    lines.append("")
    for row in association.itertuples():
        lines.append(
            f"- {feature_names[row.feature]}: chi-square={row.chi_square:.3f}, "
            f"permutation p={row.permutation_p:.4f}, Cramer's V={row.cramers_v:.3f}."
        )
    lines.extend(["", "### Multivariable Firth logistic model", ""])
    for row in firth_global.itertuples():
        lines.append(
            f"- {feature_names[row.factor]}: parameter-bootstrap global p="
            f"{row.parameter_bootstrap_p:.4f} ({row.successful_simulations} successful simulations)."
        )
    type_coefficient = firth_coefficients.loc[firth_coefficients["factor"] == "glass_type"].iloc[0]
    type_probabilities = adjusted_probabilities.loc[
        adjusted_probabilities["factor"] == "glass_type"
    ].set_index("level")
    lines.append(
        f"- 铅钡 vs 高钾 adjusted OR={type_coefficient.adjusted_odds_ratio:.3f} "
        f"(bootstrap 95% CI {type_coefficient.or_ci_low:.3f} to {type_coefficient.or_ci_high:.3f})."
    )
    for level, row in type_probabilities.iterrows():
        lines.append(
            f"- {level} adjusted weathering probability={row.adjusted_probability:.1%} "
            f"(95% CI {row.probability_ci_low:.1%} to {row.probability_ci_high:.1%})."
        )
    type_or_column = next(
        column for column in specification_sensitivity.columns
        if column.startswith("or_glass_type_")
    )
    type_or_values = specification_sensitivity[type_or_column].dropna()
    lines.append(
        f"- Specification sensitivity: adjusted type OR ranges from "
        f"{type_or_values.min():.3f} to {type_or_values.max():.3f}; therefore adjusted "
        "probabilities and bootstrap intervals are preferred over interpreting the OR as a stable physical multiplier."
    )
    lines.extend(["", "## Composition differences", ""])
    if significant.empty:
        lines.append("- No component remained significant at BH-adjusted q < 0.05 in the artifact-state permutation analysis.")
    else:
        for glass_type, subset in significant.groupby("glass_type"):
            text = ", ".join(
                f"{r.component} ({r.mean_difference:+.2f} points, q={r.bh_q:.4f})"
                for r in subset.itertuples()
            )
            lines.append(f"- {glass_type}: {text}.")
    lines.extend([
        "",
        "The component test uses one mean record per artifact and point-weathering state, reducing repeated-point weighting. Raw percentages are retained for interpretation.",
        "",
        "## Pre-weathering prediction",
        "",
        f"- Primary model: type-specific median CLR shift, pseudocount={PSEUDOCOUNT}, effect cap=+/-{EFFECT_CAP} log units.",
        f"- Predicted valid weathered sampling points: {len(predictions)}; all predicted rows close to 100%.",
        f"- Capped type-component effects: {int(effects['effect_was_capped'].sum())} of {len(effects)}.",
    ])
    for row in sensitivity.itertuples():
        lines.append(
            f"- Pseudocount {row.pseudocount:.2f}: mean absolute change versus 0.10 = "
            f"{row.mean_absolute_change_vs_0_10:.3f} percentage points; maximum = "
            f"{row.maximum_absolute_change_vs_0_10:.3f}."
        )
    if not paired_validation.empty:
        lines.append("- Paired-point check: " + "; ".join(
            f"artifact {row.artifact_id}, MAE {row.current_mae:.3f} -> {row.predicted_mae:.3f}"
            for row in paired_validation.itertuples()
        ) + ".")
    lines.extend([
        "",
        "This is a group-level counterfactual correction rather than a direct before-after supervised model. Only artifacts with both point states could support true paired calibration, so paired information is reported separately and should be used as a robustness check rather than the sole estimator.",
        "",
    ])
    return "\n".join(lines)


def main() -> None:
    script_path = Path(__file__).resolve()
    question_root = script_path.parents[1]
    project_root = script_path.parents[2]
    processed_dir = question_root / "data" / "processed"
    results_dir = question_root / "results"
    figures_dir = question_root / "figures"
    docs_dir = question_root / "docs"
    for directory in (processed_dir, results_dir, figures_dir, docs_dir):
        directory.mkdir(parents=True, exist_ok=True)

    source = find_source_workbook(project_root)
    artifacts, samples = read_workbook(source)
    valid_samples = samples.loc[samples["is_valid"]].copy()
    units = valid_samples.groupby(
        ["artifact_id", "glass_type", "pattern", "color", "point_weathering"],
        as_index=False,
    )[COMPONENTS].mean()

    rng = np.random.default_rng(SEED)
    association_rows = []
    for feature in ("glass_type", "pattern", "color"):
        result = association_test(artifacts[feature], artifacts["surface_weathering"], rng)
        association_rows.append({"feature": feature, **result})
    association = pd.DataFrame(association_rows)

    X, y, design_names, design_metadata = build_weathering_design(artifacts)
    firth_global, full_firth_fit = firth_global_factor_tests(X, y, design_metadata, rng)
    bootstrap_betas, firth_diagnostics = bootstrap_firth_coefficients(
        X, full_firth_fit, rng
    )
    firth_coefficients, adjusted_probabilities = firth_result_tables(
        X, design_names, design_metadata, full_firth_fit, bootstrap_betas
    )
    specification_sensitivity = firth_specification_sensitivity(
        X, y, design_names, design_metadata
    )

    rates = []
    for feature in ("glass_type", "pattern", "color"):
        grouped = artifacts.groupby(feature, dropna=False)["surface_weathering"]
        for level, values in grouped:
            rates.append({
                "feature": feature,
                "level": level,
                "n": len(values),
                "weathered_n": int(values.eq("风化").sum()),
                "weathering_rate": values.eq("风化").mean(),
            })
    weathering_rates = pd.DataFrame(rates)
    stats = component_statistics(units, rng)
    predictions, effects = predict_preweathering(valid_samples, units, PSEUDOCOUNT)
    sensitivity = sensitivity_analysis(valid_samples, units)
    paired_validation = validate_paired_artifacts(units, predictions)

    state_counts = units.groupby(["glass_type", "point_weathering"]).size().rename("n_artifact_states").reset_index()
    paired = units.groupby("artifact_id")["point_weathering"].nunique()
    paired_ids = paired[paired > 1].index
    paired_points = units.loc[units["artifact_id"].isin(paired_ids)].copy()

    sample_export = samples[[
        "sample_id", "artifact_id", "glass_type", "pattern", "color",
        "surface_weathering", "point_weathering", "component_sum", "is_valid", *COMPONENTS,
    ]]
    sample_export.to_csv(processed_dir / "cleaned-sampling-points.csv", index=False, encoding="utf-8-sig")
    units.to_csv(processed_dir / "artifact-state-compositions.csv", index=False, encoding="utf-8-sig")
    association.to_csv(results_dir / "weathering-association-tests.csv", index=False, encoding="utf-8-sig")
    firth_global.to_csv(results_dir / "firth-global-tests.csv", index=False, encoding="utf-8-sig")
    firth_coefficients.to_csv(results_dir / "firth-coefficients.csv", index=False, encoding="utf-8-sig")
    adjusted_probabilities.to_csv(
        results_dir / "firth-adjusted-probabilities.csv", index=False, encoding="utf-8-sig"
    )
    firth_diagnostics.to_csv(
        results_dir / "firth-bootstrap-diagnostics.csv", index=False, encoding="utf-8-sig"
    )
    specification_sensitivity.to_csv(
        results_dir / "firth-specification-sensitivity.csv", index=False, encoding="utf-8-sig"
    )
    weathering_rates.to_csv(results_dir / "weathering-rates.csv", index=False, encoding="utf-8-sig")
    stats.to_csv(results_dir / "component-difference-tests.csv", index=False, encoding="utf-8-sig")
    predictions.to_csv(results_dir / "preweathering-predictions.csv", index=False, encoding="utf-8-sig")
    effects.to_csv(results_dir / "weathering-clr-effects.csv", index=False, encoding="utf-8-sig")
    sensitivity.to_csv(results_dir / "prediction-sensitivity.csv", index=False, encoding="utf-8-sig")
    paired_validation.to_csv(results_dir / "paired-prediction-validation.csv", index=False, encoding="utf-8-sig")
    state_counts.to_csv(results_dir / "analysis-unit-counts.csv", index=False, encoding="utf-8-sig")
    paired_points.to_csv(results_dir / "paired-artifact-states.csv", index=False, encoding="utf-8-sig")

    write_association_svg(artifacts, figures_dir / "weathering-rates.svg")
    write_adjusted_probability_svg(
        adjusted_probabilities, figures_dir / "adjusted-weathering-probabilities.svg"
    )
    write_component_heatmap(stats, figures_dir / "component-differences.svg")
    write_prediction_svg(predictions, figures_dir / "preweathering-comparison.svg")
    report = build_report(
        artifacts, samples, association, firth_global, firth_coefficients,
        adjusted_probabilities, specification_sensitivity, stats, predictions, effects,
        sensitivity, paired_validation
    )
    (docs_dir / "preliminary-findings.md").write_text(report, encoding="utf-8")

    print(f"Source: {source}")
    print(f"Valid sampling points: {len(valid_samples)}/{len(samples)}")
    print("Artifact-state counts:")
    print(state_counts.to_string(index=False))
    print(f"Artifacts with both point-weathering states: {', '.join(paired_ids) if len(paired_ids) else 'none'}")
    print("Association tests:")
    print(association.to_string(index=False))
    print("Firth global factor tests:")
    print(firth_global.to_string(index=False))
    print("Firth type effect:")
    print(firth_coefficients.loc[firth_coefficients["factor"] == "glass_type"].to_string(index=False))
    print("BH-significant component differences:")
    print(stats.loc[stats['bh_q'] < 0.05, ["glass_type", "component", "mean_difference", "bh_q"]].to_string(index=False))
    print("Sensitivity:")
    print(sensitivity.to_string(index=False))
    print(f"Outputs: {question_root}")


if __name__ == "__main__":
    main()
