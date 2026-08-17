#!/usr/bin/env python3
"""Evaluate the fixed heliostat field for question 1."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import math
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.spatial import cKDTree
from scipy.stats import qmc


MONTH_DAYS = [306, 337, 0, 31, 61, 92, 122, 153, 184, 214, 245, 275]
SOLAR_TIMES = [9.0, 10.5, 12.0, 13.5, 15.0]
LATITUDE = math.radians(39.4)
SOLAR_RADIUS = math.radians(0.266)
MIRROR_AREA = 36.0
MIRROR_HALF_SIZE = 3.0
MIRROR_CENTER_Z = 4.0
RECEIVER_CENTER = np.array([0.0, 0.0, 80.0])
RECEIVER_RADIUS = 3.5
RECEIVER_Z_MIN = 76.0
RECEIVER_Z_MAX = 84.0
REFLECTIVITY = 0.92
EPS = 1.0e-9


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--samples", type=int, default=256)
    parser.add_argument("--seed", type=int, default=202308)
    parser.add_argument("--neighbor-radius", type=float, default=45.0)
    return parser.parse_args()


def setup_logging(output: Path) -> logging.Logger:
    (output / "logs").mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("q1-evaluation")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    formatter = logging.Formatter("%(asctime)s %(levelname)s %(message)s")
    file_handler = logging.FileHandler(output / "logs" / "run.log", encoding="utf-8")
    file_handler.setFormatter(formatter)
    stream_handler = logging.StreamHandler()
    stream_handler.setFormatter(formatter)
    logger.addHandler(file_handler)
    logger.addHandler(stream_handler)
    return logger


def read_centers(path: Path) -> np.ndarray:
    workbook = load_workbook(path, data_only=True, read_only=True)
    rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
    xy = np.asarray(rows, dtype=float)
    if xy.ndim != 2 or xy.shape[1] != 2 or not np.isfinite(xy).all():
        raise ValueError("The attachment must contain two finite numeric columns.")
    return np.column_stack([xy, np.full(len(xy), MIRROR_CENTER_Z)])


def sobol_samples(count: int, seed: int) -> np.ndarray:
    exponent = int(round(math.log2(count)))
    if count <= 0 or 2**exponent != count:
        raise ValueError("--samples must be a positive power of two")
    return qmc.Sobol(d=4, scramble=True, seed=seed).random_base2(exponent)


def solar_state(day: int, solar_time: float) -> tuple[float, np.ndarray, float]:
    delta = math.asin(math.sin(2.0 * math.pi * day / 365.0) * math.sin(math.radians(23.45)))
    omega = math.pi * (solar_time - 12.0) / 12.0
    sin_alpha = (
        math.cos(delta) * math.cos(LATITUDE) * math.cos(omega)
        + math.sin(delta) * math.sin(LATITUDE)
    )
    alpha = math.asin(np.clip(sin_alpha, -1.0, 1.0))
    direction = np.array(
        [
            -math.cos(delta) * math.sin(omega),
            math.cos(LATITUDE) * math.sin(delta)
            - math.sin(LATITUDE) * math.cos(delta) * math.cos(omega),
            sin_alpha,
        ],
        dtype=float,
    )
    direction /= np.linalg.norm(direction)
    dni = 1.366 * (0.34981 + 0.5783875 * math.exp(-0.275745 / sin_alpha))
    return alpha, direction, dni


def solar_disk_directions(sun: np.ndarray, uv: np.ndarray) -> np.ndarray:
    reference = np.array([0.0, 0.0, 1.0])
    e1 = np.cross(sun, reference)
    if np.linalg.norm(e1) < 1.0e-10:
        e1 = np.cross(sun, np.array([0.0, 1.0, 0.0]))
    e1 /= np.linalg.norm(e1)
    e2 = np.cross(sun, e1)
    e2 /= np.linalg.norm(e2)
    rho = SOLAR_RADIUS * np.sqrt(uv[:, 2])
    phi = 2.0 * math.pi * uv[:, 3]
    transverse = np.cos(phi)[:, None] * e1 + np.sin(phi)[:, None] * e2
    directions = np.cos(rho)[:, None] * sun + np.sin(rho)[:, None] * transverse
    return directions / np.linalg.norm(directions, axis=1)[:, None]


def mirror_frames(centers: np.ndarray, sun: np.ndarray) -> tuple[np.ndarray, ...]:
    target = RECEIVER_CENTER - centers
    distances = np.linalg.norm(target, axis=1)
    target /= distances[:, None]
    normals = sun + target
    normals /= np.linalg.norm(normals, axis=1)[:, None]
    vertical = np.array([0.0, 0.0, 1.0])
    u = np.cross(vertical, normals)
    u /= np.linalg.norm(u, axis=1)[:, None]
    v = np.cross(normals, u)
    v /= np.linalg.norm(v, axis=1)[:, None]
    return target, distances, normals, u, v


def ray_hits_mirrors(
    points: np.ndarray,
    directions: np.ndarray,
    candidate_ids: np.ndarray,
    centers: np.ndarray,
    normals: np.ndarray,
    axes_u: np.ndarray,
    axes_v: np.ndarray,
) -> np.ndarray:
    if candidate_ids.size == 0:
        return np.zeros(len(points), dtype=bool)
    c = centers[candidate_ids]
    n = normals[candidate_ids]
    u = axes_u[candidate_ids]
    v = axes_v[candidate_ids]
    denominator = directions @ n.T
    numerator = np.sum(c * n, axis=1)[None, :] - points @ n.T
    valid = np.abs(denominator) > EPS
    lam = np.divide(numerator, denominator, out=np.full_like(numerator, np.nan), where=valid)
    valid &= lam > 1.0e-7
    local_u = points @ u.T + lam * (directions @ u.T) - np.sum(c * u, axis=1)[None, :]
    local_v = points @ v.T + lam * (directions @ v.T) - np.sum(c * v, axis=1)[None, :]
    valid &= np.abs(local_u) <= MIRROR_HALF_SIZE + 1.0e-8
    valid &= np.abs(local_v) <= MIRROR_HALF_SIZE + 1.0e-8
    return np.any(valid, axis=1)


def ray_hits_receiver(points: np.ndarray, directions: np.ndarray) -> np.ndarray:
    dx = directions[:, 0]
    dy = directions[:, 1]
    px = points[:, 0]
    py = points[:, 1]
    a = dx * dx + dy * dy
    b = 2.0 * (px * dx + py * dy)
    c = px * px + py * py - RECEIVER_RADIUS**2
    discriminant = b * b - 4.0 * a * c
    hit = discriminant >= 0.0
    sqrt_disc = np.sqrt(np.maximum(discriminant, 0.0))
    root1 = (-b - sqrt_disc) / (2.0 * a)
    root2 = (-b + sqrt_disc) / (2.0 * a)
    lam = np.where((root1 > EPS) & ((root1 <= root2) | (root2 <= EPS)), root1, root2)
    hit &= lam > EPS
    z = points[:, 2] + lam * directions[:, 2]
    hit &= (z >= RECEIVER_Z_MIN) & (z <= RECEIVER_Z_MAX)
    x = px + lam * dx
    y = py + lam * dy
    hit &= (directions[:, 0] * x + directions[:, 1] * y) < 0.0
    return hit


def required_neighbor_radius(sun: np.ndarray, target: np.ndarray) -> float:
    lower_shadow_z = sun[2] * math.cos(SOLAR_RADIUS) - math.sqrt(1.0 - sun[2] ** 2) * math.sin(SOLAR_RADIUS)
    target_z = target[:, 2]
    lower_reflected_z = target_z * math.cos(SOLAR_RADIUS) - np.sqrt(1.0 - target_z**2) * math.sin(SOLAR_RADIUS)
    minimum_z = min(lower_shadow_z, float(np.min(lower_reflected_z)))
    if minimum_z <= 0.0:
        raise ValueError("A geometric ray can travel downward; fixed-radius neighbor proof is invalid.")
    maximum_slope = math.sqrt(1.0 - minimum_z**2) / minimum_z
    center_offset_allowance = 2.0 * math.sqrt(2.0) * MIRROR_HALF_SIZE
    return center_offset_allowance + 6.0 * maximum_slope


def evaluate_time(
    centers: np.ndarray,
    neighbors: list[np.ndarray],
    uv: np.ndarray,
    sun: np.ndarray,
    dni: float,
    neighbor_radius: float,
) -> tuple[dict[str, float], np.ndarray]:
    target, distances, normals, axes_u, axes_v = mirror_frames(centers, sun)
    required_radius = required_neighbor_radius(sun, target)
    if required_radius > neighbor_radius + 1.0e-9:
        raise ValueError(
            f"Neighbor radius {neighbor_radius:.3f} m is below conservative bound {required_radius:.3f} m"
        )
    disk_directions = solar_disk_directions(sun, uv)
    xi = 6.0 * (uv[:, 0] - 0.5)
    zeta = 6.0 * (uv[:, 1] - 0.5)
    count = len(centers)
    eta_sb = np.empty(count)
    eta_trunc = np.empty(count)
    for i in range(count):
        points = centers[i] + xi[:, None] * axes_u[i] + zeta[:, None] * axes_v[i]
        reflected = -disk_directions + 2.0 * (disk_directions @ normals[i])[:, None] * normals[i]
        reflected /= np.linalg.norm(reflected, axis=1)[:, None]
        shadow = ray_hits_mirrors(
            points, disk_directions, neighbors[i], centers, normals, axes_u, axes_v
        )
        blocked = ray_hits_mirrors(
            points, reflected, neighbors[i], centers, normals, axes_u, axes_v
        )
        clear = ~(shadow | blocked)
        receiver_hit = ray_hits_receiver(points, reflected)
        eta_sb[i] = np.mean(clear)
        eta_trunc[i] = np.mean(receiver_hit[clear]) if np.any(clear) else 0.0
    eta_cos = normals @ sun
    eta_at = 0.99321 - 0.0001176 * distances + 1.97e-8 * distances**2
    eta_opt = eta_sb * eta_cos * eta_at * eta_trunc * REFLECTIVITY
    field_power_kw = dni * MIRROR_AREA * np.sum(eta_opt)
    metrics = {
        "dni_kw_m2": dni,
        "avg_optical_efficiency": float(np.mean(eta_opt)),
        "avg_cosine_efficiency": float(np.mean(eta_cos)),
        "avg_shadow_blocking_efficiency": float(np.mean(eta_sb)),
        "avg_truncation_efficiency": float(np.mean(eta_trunc)),
        "field_power_mw": field_power_kw / 1000.0,
        "unit_area_power_kw_m2": field_power_kw / (MIRROR_AREA * count),
        "required_neighbor_radius_m": required_radius,
    }
    mirror_metrics = np.column_stack([eta_opt, eta_cos, eta_at, eta_sb, eta_trunc])
    return metrics, mirror_metrics


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def mean_columns(rows: list[dict[str, object]], keys: list[str]) -> dict[str, float]:
    return {key: float(np.mean([float(row[key]) for row in rows])) for key in keys}


def write_summary(output: Path, monthly: list[dict[str, object]], annual: dict[str, object]) -> None:
    lines = [
        "# 问题一计算结果",
        "",
        "## 每月 21 日平均指标",
        "",
        "| 月份 | 平均光学效率 | 平均余弦效率 | 平均阴影遮挡效率 | 平均截断效率 | 单位面积输出热功率 (kW/m2) |",
        "|---:|---:|---:|---:|---:|---:|",
    ]
    for row in monthly:
        lines.append(
            f"| {row['month']} | {row['avg_optical_efficiency']:.4f} | "
            f"{row['avg_cosine_efficiency']:.4f} | {row['avg_shadow_blocking_efficiency']:.4f} | "
            f"{row['avg_truncation_efficiency']:.4f} | {row['unit_area_power_kw_m2']:.4f} |"
        )
    lines.extend(
        [
            "",
            "## 年平均指标",
            "",
            "| 年平均光学效率 | 年平均余弦效率 | 年平均阴影遮挡效率 | 年平均截断效率 | 年平均输出热功率 (MW) | 单位面积年平均输出热功率 (kW/m2) |",
            "|---:|---:|---:|---:|---:|---:|",
            (
                f"| {annual['avg_optical_efficiency']:.4f} | {annual['avg_cosine_efficiency']:.4f} | "
                f"{annual['avg_shadow_blocking_efficiency']:.4f} | {annual['avg_truncation_efficiency']:.4f} | "
                f"{annual['field_power_mw']:.4f} | {annual['unit_area_power_kw_m2']:.4f} |"
            ),
            "",
            f"- 联合射线样本数：{annual['samples_per_mirror_time']}",
            f"- Sobol 扰码种子：{annual['seed']}",
            f"- 邻镜搜索半径：{annual['neighbor_radius_m']:.1f} m",
        ]
    )
    (output / "results" / "result_summary.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_plot(output: Path, monthly: list[dict[str, object]]) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    months = [int(row["month"]) for row in monthly]
    figure, axes = plt.subplots(2, 1, figsize=(10, 8), constrained_layout=True)
    for key, label in [
        ("avg_optical_efficiency", "Optical"),
        ("avg_cosine_efficiency", "Cosine"),
        ("avg_shadow_blocking_efficiency", "Shadow/blocking"),
        ("avg_truncation_efficiency", "Truncation"),
    ]:
        axes[0].plot(months, [row[key] for row in monthly], marker="o", label=label)
    axes[0].set_ylabel("Efficiency")
    axes[0].set_ylim(0.45, 1.02)
    axes[0].set_xticks(months)
    axes[0].grid(alpha=0.25)
    axes[0].legend(ncol=2)
    axes[1].bar(months, [row["unit_area_power_kw_m2"] for row in monthly], color="#d97706")
    axes[1].set_xlabel("Month")
    axes[1].set_ylabel("Unit-area power (kW/m2)")
    axes[1].set_xticks(months)
    axes[1].grid(axis="y", alpha=0.25)
    (output / "figures").mkdir(parents=True, exist_ok=True)
    figure.savefig(output / "figures" / "monthly_performance.png", dpi=180)
    plt.close(figure)


def main() -> None:
    args = parse_args()
    for subdir in ["results", "figures", "validation", "logs"]:
        (args.output / subdir).mkdir(parents=True, exist_ok=True)
    logger = setup_logging(args.output)
    centers = read_centers(args.input)
    if len(centers) != 1745:
        raise ValueError(f"Expected 1745 heliostats, found {len(centers)}")
    uv = sobol_samples(args.samples, args.seed)
    tree = cKDTree(centers[:, :2])
    neighbor_lists = []
    for i, ids in enumerate(tree.query_ball_point(centers[:, :2], args.neighbor_radius)):
        neighbor_lists.append(np.asarray([j for j in ids if j != i], dtype=int))
    logger.info(
        "Loaded %d heliostats; samples=%d; neighbor count min/mean/max=%d/%.1f/%d",
        len(centers),
        args.samples,
        min(map(len, neighbor_lists)),
        np.mean(list(map(len, neighbor_lists))),
        max(map(len, neighbor_lists)),
    )
    time_rows: list[dict[str, object]] = []
    mirror_rows: list[dict[str, object]] = []
    metric_keys = [
        "avg_optical_efficiency",
        "avg_cosine_efficiency",
        "avg_shadow_blocking_efficiency",
        "avg_truncation_efficiency",
        "field_power_mw",
        "unit_area_power_kw_m2",
    ]
    for month, day in enumerate(MONTH_DAYS, start=1):
        for solar_time in SOLAR_TIMES:
            alpha, sun, dni = solar_state(day, solar_time)
            metrics, per_mirror = evaluate_time(
                centers, neighbor_lists, uv, sun, dni, args.neighbor_radius
            )
            row: dict[str, object] = {
                "month": month,
                "day_offset_from_march_21": day,
                "solar_time": solar_time,
                "solar_altitude_deg": math.degrees(alpha),
                **metrics,
            }
            time_rows.append(row)
            for i, values in enumerate(per_mirror, start=1):
                mirror_rows.append(
                    {
                        "month": month,
                        "solar_time": solar_time,
                        "mirror_id": i,
                        "optical_efficiency": values[0],
                        "cosine_efficiency": values[1],
                        "atmospheric_efficiency": values[2],
                        "shadow_blocking_efficiency": values[3],
                        "truncation_efficiency": values[4],
                    }
                )
        logger.info("Completed month %02d/12", month)
    monthly_rows: list[dict[str, object]] = []
    for month in range(1, 13):
        selected = [row for row in time_rows if row["month"] == month]
        monthly_rows.append({"month": month, **mean_columns(selected, metric_keys)})
    annual: dict[str, object] = {
        **mean_columns(time_rows, metric_keys),
        "samples_per_mirror_time": args.samples,
        "seed": args.seed,
        "neighbor_radius_m": args.neighbor_radius,
    }
    write_csv(args.output / "results" / "time_metrics.csv", time_rows)
    write_csv(args.output / "results" / "mirror_time_metrics.csv", mirror_rows)
    write_csv(args.output / "results" / "monthly_metrics.csv", monthly_rows)
    write_csv(args.output / "results" / "annual_metrics.csv", [annual])
    write_summary(args.output, monthly_rows, annual)
    write_plot(args.output, monthly_rows)
    symmetry_pairs = [(0, 4), (1, 3)]
    symmetry_errors = []
    for month in range(1, 13):
        selected = [row for row in time_rows if row["month"] == month]
        for left, right in symmetry_pairs:
            a = float(selected[left]["avg_optical_efficiency"])
            b = float(selected[right]["avg_optical_efficiency"])
            symmetry_errors.append(abs(a - b))
    validation = {
        "heliostat_count": len(centers),
        "coordinate_duplicates": len(centers) - len(np.unique(centers[:, :2], axis=0)),
        "efficiencies_in_unit_interval": all(
            0.0 <= float(row[key]) <= 1.0
            for row in time_rows
            for key in metric_keys[:4]
        ),
        "minimum_solar_altitude_deg": min(float(row["solar_altitude_deg"]) for row in time_rows),
        "maximum_required_neighbor_radius_m": max(
            float(row["required_neighbor_radius_m"]) for row in time_rows
        ),
        "configured_neighbor_radius_m": args.neighbor_radius,
        "maximum_morning_afternoon_optical_difference": max(symmetry_errors),
        "time_row_count": len(time_rows),
        "mirror_time_row_count": len(mirror_rows),
        "code_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
    }
    (args.output / "validation" / "checks.json").write_text(
        json.dumps(validation, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    logger.info("Annual metrics: %s", json.dumps(annual, ensure_ascii=False))


if __name__ == "__main__":
    main()
