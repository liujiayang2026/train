#!/usr/bin/env python3
"""Shared geometry, proxy, and ray-tracing model for question 2."""

from __future__ import annotations

import csv
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import numpy as np
from openpyxl import load_workbook
from scipy.spatial import cKDTree
from scipy.stats import qmc


MONTH_DAYS = [306, 337, 0, 31, 61, 92, 122, 153, 184, 214, 245, 275]
SOLAR_TIMES = [9.0, 10.5, 12.0, 13.5, 15.0]
LATITUDE = math.radians(39.4)
SOLAR_RADIUS = math.radians(0.266)
RECEIVER_RADIUS = 3.5
RECEIVER_Z_MIN = 76.0
RECEIVER_Z_MAX = 84.0
REFLECTIVITY = 0.92
SITE_RADIUS = 350.0
EXCLUSION_RADIUS = 100.0
EPS = 1.0e-9


@dataclass(frozen=True)
class Design:
    tower_x: float
    tower_y: float
    width: float
    height: float
    center_z: float
    first_radius: float
    radial_clearance: float
    tangential_clearance: float
    max_radius: float
    phase: float

    @classmethod
    def from_array(cls, values: Iterable[float]) -> "Design":
        return cls(*map(float, values))

    def as_array(self) -> np.ndarray:
        return np.asarray(
            [
                self.tower_x,
                self.tower_y,
                self.width,
                self.height,
                self.center_z,
                self.first_radius,
                self.radial_clearance,
                self.tangential_clearance,
                self.max_radius,
                self.phase,
            ],
            dtype=float,
        )

    @property
    def area(self) -> float:
        return self.width * self.height

    @property
    def minimum_spacing(self) -> float:
        return self.width + 5.0

    @property
    def radial_spacing(self) -> float:
        return self.minimum_spacing + self.radial_clearance

    @property
    def tangential_spacing(self) -> float:
        return self.minimum_spacing + self.tangential_clearance


@dataclass
class LayoutStages:
    raw: np.ndarray
    after_exclusion: np.ndarray
    accepted: np.ndarray
    ring_ids: np.ndarray


def read_xy_workbook(path: Path) -> np.ndarray:
    workbook = load_workbook(path, data_only=True, read_only=True)
    rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
    xy = np.asarray(rows, dtype=float)
    if xy.ndim != 2 or xy.shape[1] != 2 or not np.isfinite(xy).all():
        raise ValueError("Expected two finite coordinate columns in the workbook.")
    return xy


def read_scalar_csv(path: Path, key: str) -> float:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        row = next(csv.DictReader(handle))
    return float(row[key])


def generate_layout(design: Design, return_stages: bool = False) -> np.ndarray | LayoutStages:
    if design.radial_spacing <= 0.0 or design.tangential_spacing <= 0.0:
        empty = np.empty((0, 2), dtype=float)
        return LayoutStages(empty, empty, empty, np.empty(0, dtype=int)) if return_stages else empty
    radii = np.arange(
        max(design.first_radius, EXCLUSION_RADIUS),
        design.max_radius + 0.5 * design.radial_spacing,
        design.radial_spacing,
    )
    raw_parts: list[np.ndarray] = []
    ring_parts: list[np.ndarray] = []
    for k, radius in enumerate(radii):
        if radius <= design.tangential_spacing / 2.0:
            continue
        ratio = np.clip(design.tangential_spacing / (2.0 * radius), 0.0, 1.0)
        count = max(3, int(math.floor(math.pi / math.asin(ratio))))
        offset = design.phase + (math.pi / count if k % 2 else 0.0)
        angles = offset + 2.0 * math.pi * np.arange(count) / count
        points = np.column_stack(
            [
                design.tower_x + radius * np.cos(angles),
                design.tower_y + radius * np.sin(angles),
            ]
        )
        raw_parts.append(points)
        ring_parts.append(np.full(count, k, dtype=int))
    if not raw_parts:
        empty = np.empty((0, 2), dtype=float)
        return LayoutStages(empty, empty, empty, np.empty(0, dtype=int)) if return_stages else empty
    raw = np.vstack(raw_parts)
    ring_ids = np.concatenate(ring_parts)
    distance_to_tower = np.hypot(raw[:, 0] - design.tower_x, raw[:, 1] - design.tower_y)
    exclusion_mask = distance_to_tower >= EXCLUSION_RADIUS - 1.0e-8
    after_exclusion = raw[exclusion_mask]
    after_ring_ids = ring_ids[exclusion_mask]
    site_mask = np.hypot(after_exclusion[:, 0], after_exclusion[:, 1]) <= SITE_RADIUS + 1.0e-8
    accepted = after_exclusion[site_mask]
    accepted_ring_ids = after_ring_ids[site_mask]
    if return_stages:
        return LayoutStages(raw, after_exclusion, accepted, accepted_ring_ids)
    return accepted


def layout_checks(design: Design, xy: np.ndarray) -> dict[str, float | int | bool]:
    if len(xy) < 2:
        nearest = float("nan")
    else:
        distances, _ = cKDTree(xy).query(xy, k=2)
        nearest = float(np.min(distances[:, 1]))
    max_site = float(np.max(np.hypot(xy[:, 0], xy[:, 1]))) if len(xy) else float("nan")
    min_tower = (
        float(np.min(np.hypot(xy[:, 0] - design.tower_x, xy[:, 1] - design.tower_y)))
        if len(xy)
        else float("nan")
    )
    return {
        "mirror_count": int(len(xy)),
        "minimum_center_distance_m": nearest,
        "required_minimum_center_distance_m": design.minimum_spacing,
        "maximum_site_radius_m": max_site,
        "minimum_tower_distance_m": min_tower,
        "spacing_constraint_satisfied": bool(nearest + 1.0e-7 >= design.minimum_spacing),
        "site_constraint_satisfied": bool(max_site <= SITE_RADIUS + 1.0e-7),
        "tower_exclusion_satisfied": bool(min_tower + 1.0e-7 >= EXCLUSION_RADIUS),
    }


def solar_state(day: int, solar_time: float) -> tuple[float, np.ndarray, float]:
    delta = math.asin(math.sin(2.0 * math.pi * day / 365.0) * math.sin(math.radians(23.45)))
    omega = math.pi * (solar_time - 12.0) / 12.0
    sin_alpha = (
        math.cos(delta) * math.cos(LATITUDE) * math.cos(omega)
        + math.sin(delta) * math.sin(LATITUDE)
    )
    alpha = math.asin(np.clip(sin_alpha, -1.0, 1.0))
    direction = np.array(
        [
            -math.cos(delta) * math.sin(omega),
            math.cos(LATITUDE) * math.sin(delta)
            - math.sin(LATITUDE) * math.cos(delta) * math.cos(omega),
            sin_alpha,
        ],
        dtype=float,
    )
    direction /= np.linalg.norm(direction)
    dni = 1.366 * (0.34981 + 0.5783875 * math.exp(-0.275745 / sin_alpha))
    return alpha, direction, dni


def all_states() -> list[tuple[int, float, float, np.ndarray, float]]:
    states = []
    for month, day in enumerate(MONTH_DAYS, start=1):
        for solar_time in SOLAR_TIMES:
            alpha, sun, dni = solar_state(day, solar_time)
            states.append((month, solar_time, alpha, sun, dni))
    return states


def proxy_states() -> list[tuple[int, float, float, np.ndarray, float]]:
    return [state for state in all_states() if state[1] in (9.0, 12.0, 15.0)]


def validation_states() -> list[tuple[int, float, float, np.ndarray, float]]:
    return [
        state
        for state in all_states()
        if state[0] in (1, 4, 7, 10) and state[1] in (9.0, 12.0, 15.0)
    ]


def mirror_frames(
    centers: np.ndarray, sun: np.ndarray, receiver_center: np.ndarray
) -> tuple[np.ndarray, ...]:
    target = receiver_center - centers
    distances = np.linalg.norm(target, axis=1)
    target /= distances[:, None]
    normals = sun + target
    normals /= np.linalg.norm(normals, axis=1)[:, None]
    vertical = np.array([0.0, 0.0, 1.0])
    axes_u = np.cross(vertical, normals)
    small = np.linalg.norm(axes_u, axis=1) < 1.0e-10
    if np.any(small):
        axes_u[small] = np.cross(np.array([0.0, 1.0, 0.0]), normals[small])
    axes_u /= np.linalg.norm(axes_u, axis=1)[:, None]
    axes_v = np.cross(normals, axes_u)
    axes_v /= np.linalg.norm(axes_v, axis=1)[:, None]
    return target, distances, normals, axes_u, axes_v


def proxy_metrics_from_xy(
    design: Design,
    xy: np.ndarray,
    calibration_factor: float = 1.0,
    states: list[tuple[int, float, float, np.ndarray, float]] | None = None,
) -> dict[str, float]:
    if states is None:
        states = proxy_states()
    if len(xy) == 0:
        return {
            "avg_optical_efficiency": 0.0,
            "avg_cosine_efficiency": 0.0,
            "avg_shadow_blocking_efficiency": 0.0,
            "avg_truncation_efficiency": 0.0,
            "field_power_mw": 0.0,
            "unit_area_power_kw_m2": 0.0,
            "total_area_m2": 0.0,
            "mirror_count": 0.0,
        }
    centers = np.column_stack([xy, np.full(len(xy), design.center_z)])
    receiver = np.array([design.tower_x, design.tower_y, 80.0])
    target = receiver - centers
    distance = np.linalg.norm(target, axis=1)
    target /= distance[:, None]
    eta_at = 0.99321 - 0.0001176 * distance + 1.97e-8 * distance**2
    radial = np.hypot(xy[:, 0] - design.tower_x, xy[:, 1] - design.tower_y)
    packing = design.area / (design.radial_spacing * design.tangential_spacing)
    blur = 0.70 * distance * math.tan(SOLAR_RADIUS)
    eta_trunc = np.minimum(1.0, 7.0 / (design.width + 2.0 * blur)) * np.minimum(
        1.0, 8.0 / (design.height + 2.0 * blur)
    )
    eta_trunc = np.clip(eta_trunc, 0.0, 1.0)
    optical_values = []
    cosine_values = []
    sb_values = []
    power_values = []
    radial_factor = 0.85 + 0.15 * radial / max(design.max_radius, 1.0)
    for _, _, alpha, sun, dni in states:
        normals = target + sun
        normals /= np.linalg.norm(normals, axis=1)[:, None]
        eta_cos = normals @ sun
        eta_sb = 1.0 - 0.10 * packing * radial_factor / (math.sin(alpha) + 0.18)
        eta_sb = np.clip(eta_sb, 0.70, 1.0)
        eta_opt = REFLECTIVITY * eta_cos * eta_at * eta_sb * eta_trunc
        optical_values.append(float(np.mean(eta_opt)))
        cosine_values.append(float(np.mean(eta_cos)))
        sb_values.append(float(np.mean(eta_sb)))
        power_values.append(dni * design.area * float(np.sum(eta_opt)) * calibration_factor / 1000.0)
    mean_power = float(np.mean(power_values))
    total_area = design.area * len(xy)
    return {
        "avg_optical_efficiency": float(np.mean(optical_values)),
        "avg_cosine_efficiency": float(np.mean(cosine_values)),
        "avg_shadow_blocking_efficiency": float(np.mean(sb_values)),
        "avg_truncation_efficiency": float(np.mean(eta_trunc)),
        "field_power_mw": mean_power,
        "unit_area_power_kw_m2": mean_power * 1000.0 / total_area,
        "total_area_m2": total_area,
        "mirror_count": float(len(xy)),
    }


def proxy_calibration_factor(q1_xy: np.ndarray, q1_unit_power: float) -> float:
    baseline = Design(0.0, 0.0, 6.0, 6.0, 4.0, 100.0, 0.0, 0.0, 350.0, 0.0)
    raw = proxy_metrics_from_xy(baseline, q1_xy, calibration_factor=1.0)
    return q1_unit_power / raw["unit_area_power_kw_m2"]


def sobol_samples(count: int, seed: int) -> np.ndarray:
    exponent = int(round(math.log2(count)))
    if count <= 0 or 2**exponent != count:
        raise ValueError("Sample count must be a positive power of two.")
    return qmc.Sobol(d=4, scramble=True, seed=seed).random_base2(exponent)


def solar_disk_directions(sun: np.ndarray, uv: np.ndarray) -> np.ndarray:
    reference = np.array([0.0, 0.0, 1.0])
    e1 = np.cross(sun, reference)
    if np.linalg.norm(e1) < 1.0e-10:
        e1 = np.cross(sun, np.array([0.0, 1.0, 0.0]))
    e1 /= np.linalg.norm(e1)
    e2 = np.cross(sun, e1)
    e2 /= np.linalg.norm(e2)
    rho = SOLAR_RADIUS * np.sqrt(uv[:, 2])
    phi = 2.0 * math.pi * uv[:, 3]
    transverse = np.cos(phi)[:, None] * e1 + np.sin(phi)[:, None] * e2
    directions = np.cos(rho)[:, None] * sun + np.sin(rho)[:, None] * transverse
    return directions / np.linalg.norm(directions, axis=1)[:, None]


def ray_hits_mirrors(
    points: np.ndarray,
    directions: np.ndarray,
    candidate_ids: np.ndarray,
    centers: np.ndarray,
    normals: np.ndarray,
    axes_u: np.ndarray,
    axes_v: np.ndarray,
    half_width: float,
    half_height: float,
) -> np.ndarray:
    if candidate_ids.size == 0:
        return np.zeros(len(points), dtype=bool)
    c = centers[candidate_ids]
    n = normals[candidate_ids]
    u = axes_u[candidate_ids]
    v = axes_v[candidate_ids]
    denominator = directions @ n.T
    numerator = np.sum(c * n, axis=1)[None, :] - points @ n.T
    valid = np.abs(denominator) > EPS
    lam = np.divide(numerator, denominator, out=np.full_like(numerator, np.nan), where=valid)
    valid &= lam > 1.0e-7
    local_u = points @ u.T + lam * (directions @ u.T) - np.sum(c * u, axis=1)[None, :]
    local_v = points @ v.T + lam * (directions @ v.T) - np.sum(c * v, axis=1)[None, :]
    valid &= np.abs(local_u) <= half_width + 1.0e-8
    valid &= np.abs(local_v) <= half_height + 1.0e-8
    return np.any(valid, axis=1)


def ray_hits_receiver(
    points: np.ndarray, directions: np.ndarray, tower_x: float, tower_y: float
) -> np.ndarray:
    px = points[:, 0] - tower_x
    py = points[:, 1] - tower_y
    dx = directions[:, 0]
    dy = directions[:, 1]
    a = dx * dx + dy * dy
    b = 2.0 * (px * dx + py * dy)
    c = px * px + py * py - RECEIVER_RADIUS**2
    discriminant = b * b - 4.0 * a * c
    hit = discriminant >= 0.0
    sqrt_disc = np.sqrt(np.maximum(discriminant, 0.0))
    root1 = (-b - sqrt_disc) / (2.0 * a)
    root2 = (-b + sqrt_disc) / (2.0 * a)
    lam = np.where((root1 > EPS) & ((root1 <= root2) | (root2 <= EPS)), root1, root2)
    hit &= lam > EPS
    z = points[:, 2] + lam * directions[:, 2]
    hit &= (z >= RECEIVER_Z_MIN) & (z <= RECEIVER_Z_MAX)
    x = px + lam * dx
    y = py + lam * dy
    hit &= (directions[:, 0] * x + directions[:, 1] * y) < 0.0
    return hit


def required_neighbor_radius(design: Design, sun: np.ndarray, target: np.ndarray) -> float:
    lower_shadow_z = sun[2] * math.cos(SOLAR_RADIUS) - math.sqrt(1.0 - sun[2] ** 2) * math.sin(
        SOLAR_RADIUS
    )
    target_z = target[:, 2]
    lower_reflected_z = target_z * math.cos(SOLAR_RADIUS) - np.sqrt(1.0 - target_z**2) * math.sin(
        SOLAR_RADIUS
    )
    minimum_z = min(lower_shadow_z, float(np.min(lower_reflected_z)))
    if minimum_z <= 0.0:
        raise ValueError("A sampled ray can travel downward; neighbor-radius proof is invalid.")
    maximum_slope = math.sqrt(1.0 - minimum_z**2) / minimum_z
    diagonal = math.hypot(design.width, design.height)
    return diagonal + max(design.width, design.height) * maximum_slope


def evaluate_field(
    design: Design,
    xy: np.ndarray,
    samples: int,
    seed: int,
    neighbor_radius: float = 55.0,
    states: list[tuple[int, float, float, np.ndarray, float]] | None = None,
    store_per_mirror: bool = False,
    logger=None,
) -> tuple[list[dict[str, float]], np.ndarray | None]:
    if states is None:
        states = all_states()
    centers = np.column_stack([xy, np.full(len(xy), design.center_z)])
    receiver = np.array([design.tower_x, design.tower_y, 80.0])
    uv = sobol_samples(samples, seed)
    xi = design.width * (uv[:, 0] - 0.5)
    zeta = design.height * (uv[:, 1] - 0.5)
    half_width = design.width / 2.0
    half_height = design.height / 2.0
    tree = cKDTree(xy)
    neighbors = [
        np.asarray([j for j in ids if j != i], dtype=int)
        for i, ids in enumerate(tree.query_ball_point(xy, neighbor_radius))
    ]
    rows: list[dict[str, float]] = []
    per_mirror_sum = np.zeros((len(xy), 5), dtype=float) if store_per_mirror else None
    for state_index, (month, solar_time, alpha, sun, dni) in enumerate(states, start=1):
        target, distances, normals, axes_u, axes_v = mirror_frames(centers, sun, receiver)
        required_radius = required_neighbor_radius(design, sun, target)
        if required_radius > neighbor_radius + 1.0e-9:
            raise ValueError(
                f"Neighbor radius {neighbor_radius:.3f} m is below bound {required_radius:.3f} m"
            )
        disk_directions = solar_disk_directions(sun, uv)
        eta_sb = np.empty(len(xy), dtype=float)
        eta_trunc = np.empty(len(xy), dtype=float)
        diagonal = math.hypot(design.width, design.height)
        for i in range(len(xy)):
            points = centers[i] + xi[:, None] * axes_u[i] + zeta[:, None] * axes_v[i]
            reflected = -disk_directions + 2.0 * (disk_directions @ normals[i])[:, None] * normals[i]
            reflected /= np.linalg.norm(reflected, axis=1)[:, None]
            ids = neighbors[i]
            if ids.size:
                delta = xy[ids] - xy[i]
                shadow_ids = ids[(delta @ sun[:2]) >= -diagonal]
                block_ids = ids[(delta @ target[i, :2]) >= -diagonal]
            else:
                shadow_ids = ids
                block_ids = ids
            shadow = ray_hits_mirrors(
                points,
                disk_directions,
                shadow_ids,
                centers,
                normals,
                axes_u,
                axes_v,
                half_width,
                half_height,
            )
            blocked = ray_hits_mirrors(
                points,
                reflected,
                block_ids,
                centers,
                normals,
                axes_u,
                axes_v,
                half_width,
                half_height,
            )
            clear = ~(shadow | blocked)
            receiver_hit = ray_hits_receiver(points, reflected, design.tower_x, design.tower_y)
            eta_sb[i] = float(np.mean(clear))
            eta_trunc[i] = float(np.mean(receiver_hit[clear])) if np.any(clear) else 0.0
        eta_cos = normals @ sun
        eta_at = 0.99321 - 0.0001176 * distances + 1.97e-8 * distances**2
        eta_opt = REFLECTIVITY * eta_sb * eta_cos * eta_at * eta_trunc
        power_mw = dni * design.area * float(np.sum(eta_opt)) / 1000.0
        rows.append(
            {
                "month": float(month),
                "solar_time": float(solar_time),
                "solar_altitude_deg": math.degrees(alpha),
                "dni_kw_m2": float(dni),
                "avg_optical_efficiency": float(np.mean(eta_opt)),
                "avg_cosine_efficiency": float(np.mean(eta_cos)),
                "avg_shadow_blocking_efficiency": float(np.mean(eta_sb)),
                "avg_truncation_efficiency": float(np.mean(eta_trunc)),
                "field_power_mw": power_mw,
                "unit_area_power_kw_m2": power_mw * 1000.0 / (design.area * len(xy)),
                "required_neighbor_radius_m": required_radius,
            }
        )
        if per_mirror_sum is not None:
            per_mirror_sum += np.column_stack([eta_opt, eta_cos, eta_at, eta_sb, eta_trunc])
        if logger is not None and (state_index % 5 == 0 or state_index == len(states)):
            logger.info("Ray evaluation completed state %d/%d", state_index, len(states))
    if per_mirror_sum is not None:
        per_mirror_sum /= len(states)
    return rows, per_mirror_sum


def aggregate_time_rows(
    rows: list[dict[str, float]], total_area: float
) -> tuple[list[dict[str, float]], dict[str, float]]:
    keys = [
        "avg_optical_efficiency",
        "avg_cosine_efficiency",
        "avg_shadow_blocking_efficiency",
        "avg_truncation_efficiency",
        "field_power_mw",
        "unit_area_power_kw_m2",
    ]
    monthly: list[dict[str, float]] = []
    for month in sorted({int(row["month"]) for row in rows}):
        selected = [row for row in rows if int(row["month"]) == month]
        monthly.append(
            {"month": float(month), **{key: float(np.mean([row[key] for row in selected])) for key in keys}}
        )
    annual = {key: float(np.mean([row[key] for row in rows])) for key in keys}
    annual["total_area_m2"] = float(total_area)
    return monthly, annual
