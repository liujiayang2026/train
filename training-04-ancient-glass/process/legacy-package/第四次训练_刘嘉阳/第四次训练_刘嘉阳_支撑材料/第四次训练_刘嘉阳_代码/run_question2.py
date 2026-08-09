from __future__ import annotations

import csv
import html
import math
import re
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd


SEED = 20220804
PSEUDOCOUNT = 0.10
BOOTSTRAP_ITERATIONS = 300

COMPONENTS = [
    "SiO2", "Na2O", "K2O", "CaO", "MgO", "Al2O3", "Fe2O3",
    "CuO", "PbO", "BaO", "P2O5", "SrO", "SnO2", "SO2",
]

RAW_COMPONENT_LABELS = [
    "二氧化硅(SiO2)", "氧化钠(Na2O)", "氧化钾(K2O)", "氧化钙(CaO)",
    "氧化镁(MgO)", "氧化铝(Al2O3)", "氧化铁(Fe2O3)", "氧化铜(CuO)",
    "氧化铅(PbO)", "氧化钡(BaO)", "五氧化二磷(P2O5)", "氧化锶(SrO)",
    "氧化锡(SnO2)", "二氧化硫(SO2)",
]

MAJOR_FEATURES = ["SiO2", "K2O", "CaO", "Al2O3", "PbO", "BaO"]
HIGH_K_FEATURES = ["SiO2", "K2O", "CaO", "MgO", "Al2O3", "Fe2O3", "CuO"]
LEAD_BARIUM_FEATURES = ["SiO2", "PbO", "BaO", "P2O5", "SrO", "CuO"]

TYPE_HIGH_K = "高钾"
TYPE_LEAD_BARIUM = "铅钡"


def project_root() -> Path:
    return Path(__file__).resolve().parents[2]


def ensure_directories(root: Path) -> dict[str, Path]:
    q2 = root / "question-2"
    dirs = {
        "code": q2 / "code",
        "docs": q2 / "docs",
        "processed": q2 / "data" / "processed",
        "results": q2 / "results",
        "figures": q2 / "figures",
        "qa": q2 / "qa",
    }
    for path in dirs.values():
        path.mkdir(parents=True, exist_ok=True)
    return dirs


def find_source_workbook(root: Path) -> Path:
    candidates = [
        path for path in root.rglob("*.xlsx")
        if not path.name.startswith("~$") and "question-1" not in path.parts
        and "question-2" not in path.parts
    ]
    if len(candidates) != 1:
        raise RuntimeError(f"Expected one source workbook, found {candidates}")
    return candidates[0]


def read_workbook(path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)

    artifact_rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    artifacts = pd.DataFrame(
        artifact_rows[1:],
        columns=["artifact_id", "pattern", "glass_type", "color", "surface_weathering"],
    )
    artifacts["artifact_id"] = artifacts["artifact_id"].astype(int).astype(str).str.zfill(2)
    artifacts["color"] = artifacts["color"].fillna("未知")

    sample_rows = list(workbook.worksheets[1].iter_rows(values_only=True))
    samples = pd.DataFrame(sample_rows[1:], columns=["sample_id", *COMPONENTS])
    samples["sample_id"] = samples["sample_id"].astype(str)
    samples["artifact_id"] = (
        samples["sample_id"].str.extract(r"^(\d+)", expand=False).astype(int).astype(str).str.zfill(2)
    )
    samples[COMPONENTS] = samples[COMPONENTS].apply(pd.to_numeric, errors="coerce").fillna(0.0)
    samples["component_sum"] = samples[COMPONENTS].sum(axis=1)
    samples["is_valid"] = samples["component_sum"].between(85.0, 105.0)
    samples = samples.merge(artifacts, on="artifact_id", how="left", validate="many_to_one")
    samples["point_weathering"] = samples["surface_weathering"]
    samples.loc[samples["sample_id"].str.contains("未风化点", regex=False), "point_weathering"] = "无风化"
    samples.loc[samples["sample_id"].str.contains("严重风化点", regex=False), "point_weathering"] = "风化"
    return artifacts, samples


def read_preweathering_predictions(root: Path) -> pd.DataFrame:
    path = root / "question-1" / "results" / "preweathering-predictions.csv"
    if not path.exists():
        raise FileNotFoundError(f"Missing question 1 prediction file: {path}")
    return pd.read_csv(path, encoding="utf-8-sig")


def build_corrected_samples(samples: pd.DataFrame, predictions: pd.DataFrame) -> pd.DataFrame:
    pred_map: dict[str, dict[str, float]] = {}
    for _, row in predictions.iterrows():
        sample_id = str(row["sample_id"])
        pred_map[sample_id] = {component: float(row[f"predicted_{component}"]) for component in COMPONENTS}

    rows: list[dict[str, object]] = []
    for _, row in samples[samples["is_valid"]].iterrows():
        output = {
            "sample_id": row["sample_id"],
            "artifact_id": row["artifact_id"],
            "glass_type": row["glass_type"],
            "pattern": row["pattern"],
            "color": row["color"],
            "surface_weathering": row["surface_weathering"],
            "point_weathering": row["point_weathering"],
            "component_sum": row["component_sum"],
            "source": "raw_unweathered",
        }
        if row["point_weathering"] == "风化":
            predicted = pred_map.get(str(row["sample_id"]))
            if predicted is None:
                values = {component: float(row[component]) for component in COMPONENTS}
                output["source"] = "raw_weathered_missing_prediction"
            else:
                values = predicted
                output["source"] = "question1_preweathering_prediction"
        else:
            values = {component: float(row[component]) for component in COMPONENTS}
        output.update(values)
        rows.append(output)
    return pd.DataFrame(rows)


def aggregate_artifacts(samples: pd.DataFrame) -> pd.DataFrame:
    meta = ["artifact_id", "glass_type", "pattern", "color", "surface_weathering"]
    grouped = samples.groupby(meta, as_index=False)[COMPONENTS].mean()
    grouped["sample_count"] = samples.groupby(meta)["sample_id"].count().to_numpy()
    grouped = close_compositions(grouped, COMPONENTS)
    return grouped.sort_values("artifact_id").reset_index(drop=True)


def build_raw_artifact_data(samples: pd.DataFrame) -> pd.DataFrame:
    meta = ["artifact_id", "glass_type", "pattern", "color", "surface_weathering"]
    valid = samples[samples["is_valid"]].copy()
    grouped = valid.groupby(meta, as_index=False)[COMPONENTS].mean()
    grouped["sample_count"] = valid.groupby(meta)["sample_id"].count().to_numpy()
    grouped = close_compositions(grouped, COMPONENTS)
    return grouped.sort_values("artifact_id").reset_index(drop=True)


def close_compositions(frame: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    result = frame.copy()
    sums = result[columns].sum(axis=1).replace(0, np.nan)
    result[columns] = result[columns].div(sums, axis=0).fillna(0.0) * 100.0
    result["closed_sum"] = result[columns].sum(axis=1)
    return result


def clr_transform(values: np.ndarray, pseudocount: float = PSEUDOCOUNT) -> np.ndarray:
    values = np.asarray(values, dtype=float)
    row_sums = values.sum(axis=1, keepdims=True)
    closed = np.divide(values, row_sums, out=np.zeros_like(values), where=row_sums > 0) * 100.0
    logs = np.log(closed + pseudocount)
    return logs - logs.mean(axis=1, keepdims=True)


def standardize(values: np.ndarray) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    center = values.mean(axis=0)
    scale = values.std(axis=0, ddof=0)
    scale = np.where(scale > 1e-12, scale, 1.0)
    return (values - center) / scale, center, scale


def pca_2d(values: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    values = np.asarray(values, dtype=float)
    centered = values - values.mean(axis=0)
    _, singular_values, vt = np.linalg.svd(centered, full_matrices=False)
    coords = centered @ vt[:2].T
    variance = singular_values**2
    explained = variance / variance.sum() if variance.sum() > 0 else np.zeros_like(variance)
    if coords.shape[1] < 2:
        coords = np.column_stack([coords[:, 0], np.zeros(coords.shape[0])])
        explained = np.pad(explained, (0, 2 - len(explained)))
    return coords[:, :2], explained[:2]


def pairwise_squared_distances(values: np.ndarray, centers: np.ndarray) -> np.ndarray:
    return ((values[:, None, :] - centers[None, :, :]) ** 2).sum(axis=2)


def kmeans(values: np.ndarray, k: int, rng: np.random.Generator, n_init: int = 60) -> tuple[np.ndarray, np.ndarray, float]:
    values = np.asarray(values, dtype=float)
    n = values.shape[0]
    if k > n:
        raise ValueError("k cannot exceed sample size")

    best_labels: np.ndarray | None = None
    best_centers: np.ndarray | None = None
    best_inertia = float("inf")

    for _ in range(n_init):
        first = int(rng.integers(0, n))
        centers = [values[first]]
        while len(centers) < k:
            distances = pairwise_squared_distances(values, np.array(centers)).min(axis=1)
            if distances.sum() <= 1e-12:
                remaining = [idx for idx in range(n) if idx not in set(np.argmin(pairwise_squared_distances(values, np.array(centers)), axis=1))]
                next_idx = int(rng.choice(remaining or np.arange(n)))
            else:
                probabilities = distances / distances.sum()
                next_idx = int(rng.choice(np.arange(n), p=probabilities))
            centers.append(values[next_idx])
        centers_array = np.array(centers, dtype=float)

        labels = np.zeros(n, dtype=int)
        for _iteration in range(100):
            labels = pairwise_squared_distances(values, centers_array).argmin(axis=1)
            new_centers = centers_array.copy()
            for cluster in range(k):
                mask = labels == cluster
                if mask.any():
                    new_centers[cluster] = values[mask].mean(axis=0)
                else:
                    farthest = pairwise_squared_distances(values, centers_array).min(axis=1).argmax()
                    new_centers[cluster] = values[farthest]
            if np.allclose(new_centers, centers_array, atol=1e-10):
                break
            centers_array = new_centers

        inertia = float(pairwise_squared_distances(values, centers_array).min(axis=1).sum())
        if inertia < best_inertia:
            best_inertia = inertia
            best_labels = labels.copy()
            best_centers = centers_array.copy()

    if best_labels is None or best_centers is None:
        raise RuntimeError("K-means failed")
    return best_labels, best_centers, best_inertia


def silhouette_score(values: np.ndarray, labels: np.ndarray) -> float:
    values = np.asarray(values, dtype=float)
    labels = np.asarray(labels)
    unique = np.unique(labels)
    if len(unique) < 2 or len(unique) >= len(labels):
        return 0.0
    distances = np.sqrt(((values[:, None, :] - values[None, :, :]) ** 2).sum(axis=2))
    scores = []
    for idx in range(len(labels)):
        same = labels == labels[idx]
        if same.sum() <= 1:
            a = 0.0
        else:
            a = float(distances[idx, same].sum() / (same.sum() - 1))
        other_means = [
            float(distances[idx, labels == label].mean())
            for label in unique
            if label != labels[idx] and (labels == label).any()
        ]
        b = min(other_means) if other_means else 0.0
        denom = max(a, b)
        scores.append(0.0 if denom <= 1e-12 else (b - a) / denom)
    return float(np.mean(scores))


def ward_clusters(values: np.ndarray, target_k: int) -> tuple[np.ndarray, list[dict[str, object]]]:
    values = np.asarray(values, dtype=float)
    clusters: list[list[int]] = [[idx] for idx in range(values.shape[0])]
    history: list[dict[str, object]] = []

    def centroid(cluster: list[int]) -> np.ndarray:
        return values[cluster].mean(axis=0)

    while len(clusters) > target_k:
        best_pair: tuple[int, int] | None = None
        best_cost = float("inf")
        for i in range(len(clusters)):
            for j in range(i + 1, len(clusters)):
                ci = centroid(clusters[i])
                cj = centroid(clusters[j])
                ni = len(clusters[i])
                nj = len(clusters[j])
                cost = (ni * nj / (ni + nj)) * float(((ci - cj) ** 2).sum())
                if cost < best_cost:
                    best_cost = cost
                    best_pair = (i, j)
        if best_pair is None:
            break
        i, j = best_pair
        merged = clusters[i] + clusters[j]
        history.append({
            "merge_step": len(history) + 1,
            "cluster_a": sorted(clusters[i]),
            "cluster_b": sorted(clusters[j]),
            "merged_size": len(merged),
            "ward_cost": best_cost,
        })
        clusters = [cluster for idx, cluster in enumerate(clusters) if idx not in {i, j}]
        clusters.append(merged)

    labels = np.empty(values.shape[0], dtype=int)
    clusters = sorted(clusters, key=lambda cluster: min(cluster))
    for label, cluster in enumerate(clusters, start=1):
        labels[cluster] = label
    return labels, history


def best_threshold(values: np.ndarray, labels: np.ndarray) -> dict[str, object]:
    order = np.argsort(values)
    sorted_values = values[order]
    candidates = [sorted_values[0] - 1e-9]
    candidates += [
        float((sorted_values[i] + sorted_values[i + 1]) / 2.0)
        for i in range(len(sorted_values) - 1)
    ]
    candidates.append(sorted_values[-1] + 1e-9)

    best: dict[str, object] | None = None
    for threshold in candidates:
        pred = np.where(values >= threshold, TYPE_LEAD_BARIUM, TYPE_HIGH_K)
        accuracy = float((pred == labels).mean())
        tp = int(((pred == TYPE_LEAD_BARIUM) & (labels == TYPE_LEAD_BARIUM)).sum())
        tn = int(((pred == TYPE_HIGH_K) & (labels == TYPE_HIGH_K)).sum())
        fp = int(((pred == TYPE_LEAD_BARIUM) & (labels == TYPE_HIGH_K)).sum())
        fn = int(((pred == TYPE_HIGH_K) & (labels == TYPE_LEAD_BARIUM)).sum())
        sensitivity = tp / max(1, tp + fn)
        specificity = tn / max(1, tn + fp)
        youden = sensitivity + specificity - 1.0
        current = {
            "threshold": float(threshold),
            "accuracy": accuracy,
            "sensitivity": sensitivity,
            "specificity": specificity,
            "youden": youden,
            "tp": tp,
            "tn": tn,
            "fp": fp,
            "fn": fn,
            "predictions": pred,
        }
        if best is None or (youden, accuracy) > (float(best["youden"]), float(best["accuracy"])):
            best = current
    assert best is not None
    return best


def leave_one_out_threshold(values: np.ndarray, labels: np.ndarray) -> tuple[np.ndarray, float]:
    predictions = []
    for idx in range(len(values)):
        train_mask = np.ones(len(values), dtype=bool)
        train_mask[idx] = False
        threshold = float(best_threshold(values[train_mask], labels[train_mask])["threshold"])
        predictions.append(TYPE_LEAD_BARIUM if values[idx] >= threshold else TYPE_HIGH_K)
    predictions_array = np.array(predictions, dtype=object)
    return predictions_array, float((predictions_array == labels).mean())


def nearest_centroid_cv(features: np.ndarray, labels: np.ndarray) -> tuple[np.ndarray, float]:
    predictions = []
    for idx in range(len(labels)):
        train_mask = np.ones(len(labels), dtype=bool)
        train_mask[idx] = False
        x_train, center, scale = standardize(features[train_mask])
        x_test = (features[idx] - center) / scale
        centroids = {
            label: x_train[labels[train_mask] == label].mean(axis=0)
            for label in np.unique(labels)
        }
        distances = {label: float(((x_test - centroid) ** 2).sum()) for label, centroid in centroids.items()}
        predictions.append(min(distances, key=distances.get))
    predictions_array = np.array(predictions, dtype=object)
    return predictions_array, float((predictions_array == labels).mean())


def choose_k(values: np.ndarray, max_k: int, rng: np.random.Generator) -> tuple[int, pd.DataFrame]:
    rows = []
    max_k = min(max_k, len(values) - 1)
    for k in range(2, max_k + 1):
        labels, _, inertia = kmeans(values, k, rng)
        rows.append({
            "k": k,
            "inertia": inertia,
            "silhouette": silhouette_score(values, labels),
            "min_cluster_size": int(min(np.bincount(labels, minlength=k))),
            "cluster_sizes": ";".join(map(str, np.bincount(labels, minlength=k).tolist())),
        })
    metrics = pd.DataFrame(rows)
    if metrics.empty:
        return 1, metrics
    viable = metrics[metrics["min_cluster_size"] >= 2]
    if viable.empty:
        viable = metrics
    best_row = viable.sort_values(["silhouette", "min_cluster_size"], ascending=[False, False]).iloc[0]
    return int(best_row["k"]), metrics


def name_subclass(glass_type: str, center: pd.Series) -> str:
    if glass_type == TYPE_HIGH_K:
        if center["SnO2"] >= 1.0:
            return "高硅含锡特殊型"
        if center["SiO2"] >= 85 and center["K2O"] <= 2:
            return "高硅低钾型"
        if center["K2O"] >= 4 or center["CaO"] >= 3 or center["Al2O3"] >= 4:
            return "高钾钙铝型"
        if center["CuO"] >= 2.5 or center["Fe2O3"] >= 1.5:
            return "铜铁富集型"
        return "高钾过渡型"
    if center["BaO"] >= 15 and center["CuO"] >= 2:
        return "富钡铜型"
    if center["P2O5"] >= 3 or center["SrO"] >= 0.45:
        return "磷锶富集型"
    if center["PbO"] >= center["BaO"] * 4:
        return "富铅型"
    if center["PbO"] >= 10 and center["BaO"] >= 5:
        return "铅钡均衡型"
    return "铅钡过渡型"


def subclass_analysis(
    artifacts: pd.DataFrame,
    glass_type: str,
    features: list[str],
    max_k: int,
    rng: np.random.Generator,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, list[dict[str, object]], pd.DataFrame]:
    subset = artifacts[artifacts["glass_type"] == glass_type].copy().reset_index(drop=True)
    clr = clr_transform(subset[features].to_numpy(), PSEUDOCOUNT)
    standardized, _, _ = standardize(clr)
    selected_k, metrics = choose_k(standardized, max_k, rng)
    if glass_type == TYPE_HIGH_K and 2 in set(metrics.get("k", pd.Series(dtype=int))):
        selected_k = 2
    if glass_type == TYPE_LEAD_BARIUM and 3 in set(metrics.get("k", pd.Series(dtype=int))):
        selected_k = 3
    labels, _, inertia = kmeans(standardized, selected_k, rng, n_init=100)
    ward_labels, history = ward_clusters(standardized, selected_k)
    pca_coords, explained = pca_2d(standardized)

    subset["subclass_id"] = [f"{'K' if glass_type == TYPE_HIGH_K else 'LB'}-{label + 1}" for label in labels]
    subset["cluster_number"] = labels + 1
    subset["ward_cluster"] = ward_labels
    subset["pca1"] = pca_coords[:, 0]
    subset["pca2"] = pca_coords[:, 1]
    subset["pca1_explained"] = explained[0]
    subset["pca2_explained"] = explained[1]

    centers = []
    for cluster_number in sorted(subset["cluster_number"].unique()):
        mask = subset["cluster_number"] == cluster_number
        center = subset.loc[mask, COMPONENTS].median()
        subclass_name = name_subclass(glass_type, center)
        subset.loc[mask, "subclass_name"] = subclass_name
        centers.append({
            "glass_type": glass_type,
            "cluster_number": cluster_number,
            "subclass_id": f"{'K' if glass_type == TYPE_HIGH_K else 'LB'}-{cluster_number}",
            "subclass_name": subclass_name,
            "n": int(mask.sum()),
            "artifact_ids": ";".join(subset.loc[mask, "artifact_id"].astype(str).tolist()),
            "silhouette": silhouette_score(standardized, labels),
            "inertia": inertia,
            **{component: float(center[component]) for component in COMPONENTS},
        })
    centers_frame = pd.DataFrame(centers)

    pca_frame = subset[[
        "artifact_id", "glass_type", "subclass_id", "subclass_name",
        "pca1", "pca2", "pca1_explained", "pca2_explained",
    ]].copy()
    return subset, centers_frame, metrics, history, pca_frame


def adjusted_rand_index(labels_a: np.ndarray, labels_b: np.ndarray) -> float:
    labels_a = np.asarray(labels_a)
    labels_b = np.asarray(labels_b)
    n = len(labels_a)
    if n < 2:
        return 1.0

    def comb2(x: int) -> float:
        return x * (x - 1) / 2.0

    contingency = {}
    for a, b in zip(labels_a, labels_b):
        contingency[(a, b)] = contingency.get((a, b), 0) + 1
    sum_comb = sum(comb2(v) for v in contingency.values())
    row_counts = {}
    col_counts = {}
    for a in np.unique(labels_a):
        row_counts[a] = int((labels_a == a).sum())
    for b in np.unique(labels_b):
        col_counts[b] = int((labels_b == b).sum())
    row_comb = sum(comb2(v) for v in row_counts.values())
    col_comb = sum(comb2(v) for v in col_counts.values())
    total_comb = comb2(n)
    expected = row_comb * col_comb / total_comb if total_comb else 0.0
    maximum = 0.5 * (row_comb + col_comb)
    denom = maximum - expected
    return 0.0 if abs(denom) < 1e-12 else float((sum_comb - expected) / denom)


def compare_raw_corrected(
    corrected: pd.DataFrame,
    raw: pd.DataFrame,
    assignments: pd.DataFrame,
    rng: np.random.Generator,
) -> pd.DataFrame:
    rows = []
    for glass_type, features, max_k in [
        (TYPE_HIGH_K, HIGH_K_FEATURES, 3),
        (TYPE_LEAD_BARIUM, LEAD_BARIUM_FEATURES, 4),
    ]:
        raw_subset = raw[raw["glass_type"] == glass_type].copy().reset_index(drop=True)
        corrected_subset = corrected[corrected["glass_type"] == glass_type].copy().reset_index(drop=True)
        common = sorted(set(raw_subset["artifact_id"]) & set(corrected_subset["artifact_id"]))
        raw_subset = raw_subset[raw_subset["artifact_id"].isin(common)].sort_values("artifact_id").reset_index(drop=True)
        corrected_subset = corrected_subset[corrected_subset["artifact_id"].isin(common)].sort_values("artifact_id").reset_index(drop=True)
        base = assignments[assignments["glass_type"] == glass_type].set_index("artifact_id").loc[common]
        k = base["cluster_number"].nunique()
        raw_values, _, _ = standardize(clr_transform(raw_subset[features].to_numpy(), PSEUDOCOUNT))
        raw_labels, _, _ = kmeans(raw_values, k, rng, n_init=100)
        rows.append({
            "glass_type": glass_type,
            "comparison": "raw_valid_vs_weathering_corrected",
            "common_artifacts": len(common),
            "k": k,
            "adjusted_rand_index": adjusted_rand_index(base["cluster_number"].to_numpy(), raw_labels + 1),
        })
    return pd.DataFrame(rows)


def pseudocount_sensitivity(
    artifacts: pd.DataFrame,
    assignments: pd.DataFrame,
    rng: np.random.Generator,
) -> pd.DataFrame:
    rows = []
    for glass_type, features in [
        (TYPE_HIGH_K, HIGH_K_FEATURES),
        (TYPE_LEAD_BARIUM, LEAD_BARIUM_FEATURES),
    ]:
        subset = artifacts[artifacts["glass_type"] == glass_type].copy().sort_values("artifact_id").reset_index(drop=True)
        base = assignments[assignments["glass_type"] == glass_type].copy().sort_values("artifact_id")
        k = base["cluster_number"].nunique()
        for pseudocount in [0.05, 0.10, 0.50]:
            values, _, _ = standardize(clr_transform(subset[features].to_numpy(), pseudocount))
            labels, _, _ = kmeans(values, k, rng, n_init=100)
            rows.append({
                "glass_type": glass_type,
                "pseudocount": pseudocount,
                "k": k,
                "adjusted_rand_index_vs_0_10": adjusted_rand_index(base["cluster_number"].to_numpy(), labels + 1),
            })
    return pd.DataFrame(rows)


def feature_sensitivity(
    artifacts: pd.DataFrame,
    assignments: pd.DataFrame,
    rng: np.random.Generator,
) -> pd.DataFrame:
    rows = []
    settings = [
        (TYPE_HIGH_K, HIGH_K_FEATURES, ["K2O", "CaO", "Al2O3"]),
        (TYPE_LEAD_BARIUM, LEAD_BARIUM_FEATURES, ["PbO", "BaO", "P2O5"]),
    ]
    for glass_type, features, drops in settings:
        subset = artifacts[artifacts["glass_type"] == glass_type].copy().sort_values("artifact_id").reset_index(drop=True)
        base = assignments[assignments["glass_type"] == glass_type].copy().sort_values("artifact_id")
        k = base["cluster_number"].nunique()
        for dropped in drops:
            kept = [feature for feature in features if feature != dropped]
            values, _, _ = standardize(clr_transform(subset[kept].to_numpy(), PSEUDOCOUNT))
            labels, _, _ = kmeans(values, k, rng, n_init=100)
            rows.append({
                "glass_type": glass_type,
                "dropped_feature": dropped,
                "remaining_features": ";".join(kept),
                "adjusted_rand_index": adjusted_rand_index(base["cluster_number"].to_numpy(), labels + 1),
            })
    return pd.DataFrame(rows)


def bootstrap_stability(
    artifacts: pd.DataFrame,
    assignments: pd.DataFrame,
    rng: np.random.Generator,
) -> pd.DataFrame:
    rows = []
    settings = [
        (TYPE_HIGH_K, HIGH_K_FEATURES),
        (TYPE_LEAD_BARIUM, LEAD_BARIUM_FEATURES),
    ]
    for glass_type, features in settings:
        subset = artifacts[artifacts["glass_type"] == glass_type].copy().sort_values("artifact_id").reset_index(drop=True)
        base = assignments[assignments["glass_type"] == glass_type].copy().set_index("artifact_id")
        k = base["cluster_number"].nunique()
        values, _, _ = standardize(clr_transform(subset[features].to_numpy(), PSEUDOCOUNT))
        artifact_ids = subset["artifact_id"].to_numpy()
        base_labels = base.loc[artifact_ids, "cluster_number"].to_numpy(dtype=int)
        base_centroids = {
            cluster: values[base_labels == cluster].mean(axis=0)
            for cluster in sorted(np.unique(base_labels))
        }
        stable_counts = {artifact_id: 0 for artifact_id in artifact_ids}
        seen_counts = {artifact_id: 0 for artifact_id in artifact_ids}

        for _ in range(BOOTSTRAP_ITERATIONS):
            draw = rng.integers(0, len(subset), len(subset))
            unique_draw = np.unique(draw)
            if len(unique_draw) < k:
                continue
            boot_values = values[unique_draw]
            boot_ids = artifact_ids[unique_draw]
            labels, centers, _ = kmeans(boot_values, k, rng, n_init=40)
            label_map: dict[int, int] = {}
            for local_label in range(k):
                distances = {
                    base_cluster: float(((centers[local_label] - centroid) ** 2).sum())
                    for base_cluster, centroid in base_centroids.items()
                }
                label_map[local_label + 1] = min(distances, key=distances.get)
            for artifact_id, label in zip(boot_ids, labels + 1):
                seen_counts[artifact_id] += 1
                original = int(base.loc[artifact_id, "cluster_number"])
                if int(label_map[int(label)]) == original:
                    stable_counts[artifact_id] += 1

        for artifact_id in artifact_ids:
            seen = seen_counts[artifact_id]
            rows.append({
                "artifact_id": artifact_id,
                "glass_type": glass_type,
                "subclass_id": base.loc[artifact_id, "subclass_id"],
                "subclass_name": base.loc[artifact_id, "subclass_name"],
                "bootstrap_seen": seen,
                "same_numeric_cluster_rate": stable_counts[artifact_id] / seen if seen else np.nan,
            })
    return pd.DataFrame(rows)


def major_classification(artifacts: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    values = np.log((artifacts["PbO"].to_numpy() + artifacts["BaO"].to_numpy() + PSEUDOCOUNT) /
                    (artifacts["K2O"].to_numpy() + PSEUDOCOUNT))
    labels = artifacts["glass_type"].to_numpy()
    best = best_threshold(values, labels)
    loo_pred, loo_accuracy = leave_one_out_threshold(values, labels)

    clr = clr_transform(artifacts[MAJOR_FEATURES].to_numpy(), PSEUDOCOUNT)
    centroid_pred, centroid_accuracy = nearest_centroid_cv(clr, labels)

    result = artifacts[["artifact_id", "glass_type", *MAJOR_FEATURES]].copy()
    result["pb_ba_sum"] = result["PbO"] + result["BaO"]
    result["log_ratio_pbba_to_k"] = values
    result["threshold_prediction"] = best["predictions"]
    result["threshold_correct"] = result["threshold_prediction"] == result["glass_type"]
    result["loo_threshold_prediction"] = loo_pred
    result["loo_threshold_correct"] = result["loo_threshold_prediction"] == result["glass_type"]
    result["centroid_loo_prediction"] = centroid_pred
    result["centroid_loo_correct"] = result["centroid_loo_prediction"] == result["glass_type"]

    summary = pd.DataFrame([
        {
            "model": "log_ratio_threshold_in_sample",
            "threshold": best["threshold"],
            "accuracy": best["accuracy"],
            "sensitivity_lead_barium": best["sensitivity"],
            "specificity_high_k": best["specificity"],
            "youden": best["youden"],
            "tp": best["tp"],
            "tn": best["tn"],
            "fp": best["fp"],
            "fn": best["fn"],
        },
        {
            "model": "log_ratio_threshold_leave_one_out",
            "threshold": np.nan,
            "accuracy": loo_accuracy,
            "sensitivity_lead_barium": float(((loo_pred == TYPE_LEAD_BARIUM) & (labels == TYPE_LEAD_BARIUM)).sum() / max(1, (labels == TYPE_LEAD_BARIUM).sum())),
            "specificity_high_k": float(((loo_pred == TYPE_HIGH_K) & (labels == TYPE_HIGH_K)).sum() / max(1, (labels == TYPE_HIGH_K).sum())),
            "youden": np.nan,
            "tp": int(((loo_pred == TYPE_LEAD_BARIUM) & (labels == TYPE_LEAD_BARIUM)).sum()),
            "tn": int(((loo_pred == TYPE_HIGH_K) & (labels == TYPE_HIGH_K)).sum()),
            "fp": int(((loo_pred == TYPE_LEAD_BARIUM) & (labels == TYPE_HIGH_K)).sum()),
            "fn": int(((loo_pred == TYPE_HIGH_K) & (labels == TYPE_LEAD_BARIUM)).sum()),
        },
        {
            "model": "clr_nearest_centroid_leave_one_out",
            "threshold": np.nan,
            "accuracy": centroid_accuracy,
            "sensitivity_lead_barium": float(((centroid_pred == TYPE_LEAD_BARIUM) & (labels == TYPE_LEAD_BARIUM)).sum() / max(1, (labels == TYPE_LEAD_BARIUM).sum())),
            "specificity_high_k": float(((centroid_pred == TYPE_HIGH_K) & (labels == TYPE_HIGH_K)).sum() / max(1, (labels == TYPE_HIGH_K).sum())),
            "youden": np.nan,
            "tp": int(((centroid_pred == TYPE_LEAD_BARIUM) & (labels == TYPE_LEAD_BARIUM)).sum()),
            "tn": int(((centroid_pred == TYPE_HIGH_K) & (labels == TYPE_HIGH_K)).sum()),
            "fp": int(((centroid_pred == TYPE_LEAD_BARIUM) & (labels == TYPE_HIGH_K)).sum()),
            "fn": int(((centroid_pred == TYPE_HIGH_K) & (labels == TYPE_LEAD_BARIUM)).sum()),
        },
    ])
    return result, summary


def component_summary(artifacts: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for glass_type, group in artifacts.groupby("glass_type"):
        for component in COMPONENTS:
            rows.append({
                "glass_type": glass_type,
                "component": component,
                "n": len(group),
                "mean": group[component].mean(),
                "median": group[component].median(),
                "q25": group[component].quantile(0.25),
                "q75": group[component].quantile(0.75),
            })
    return pd.DataFrame(rows)


def svg_scatter(path: Path, points: pd.DataFrame, title: str) -> None:
    width, height = 760, 520
    margin_left, margin_right, margin_top, margin_bottom = 70, 180, 55, 70
    plot_w = width - margin_left - margin_right
    plot_h = height - margin_top - margin_bottom
    x = points["pca1"].to_numpy(dtype=float)
    y = points["pca2"].to_numpy(dtype=float)
    x_min, x_max = float(x.min()), float(x.max())
    y_min, y_max = float(y.min()), float(y.max())
    if abs(x_max - x_min) < 1e-9:
        x_min -= 1
        x_max += 1
    if abs(y_max - y_min) < 1e-9:
        y_min -= 1
        y_max += 1
    x_pad = 0.08 * (x_max - x_min)
    y_pad = 0.08 * (y_max - y_min)
    x_min -= x_pad
    x_max += x_pad
    y_min -= y_pad
    y_max += y_pad

    colors = ["#2f6f9f", "#c66b2b", "#3d8b5f", "#8b5fa8", "#b54646"]
    subclasses = sorted(points["subclass_id"].unique())
    color_map = {subclass: colors[idx % len(colors)] for idx, subclass in enumerate(subclasses)}

    def sx(value: float) -> float:
        return margin_left + (value - x_min) / (x_max - x_min) * plot_w

    def sy(value: float) -> float:
        return margin_top + (y_max - value) / (y_max - y_min) * plot_h

    elements = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<rect width="100%" height="100%" fill="#ffffff"/>',
        f'<text x="{width/2}" y="28" text-anchor="middle" font-family="Arial, sans-serif" font-size="18" font-weight="700">{html.escape(title)}</text>',
        f'<line x1="{margin_left}" y1="{margin_top+plot_h}" x2="{margin_left+plot_w}" y2="{margin_top+plot_h}" stroke="#222" stroke-width="1"/>',
        f'<line x1="{margin_left}" y1="{margin_top}" x2="{margin_left}" y2="{margin_top+plot_h}" stroke="#222" stroke-width="1"/>',
    ]
    for frac in np.linspace(0, 1, 5):
        gx = margin_left + frac * plot_w
        gy = margin_top + frac * plot_h
        elements.append(f'<line x1="{gx:.1f}" y1="{margin_top}" x2="{gx:.1f}" y2="{margin_top+plot_h}" stroke="#e5e7eb" stroke-width="1"/>')
        elements.append(f'<line x1="{margin_left}" y1="{gy:.1f}" x2="{margin_left+plot_w}" y2="{gy:.1f}" stroke="#e5e7eb" stroke-width="1"/>')

    for _, row in points.iterrows():
        cx = sx(float(row["pca1"]))
        cy = sy(float(row["pca2"]))
        color = color_map[row["subclass_id"]]
        label = html.escape(str(row["artifact_id"]))
        elements.append(f'<circle cx="{cx:.1f}" cy="{cy:.1f}" r="6" fill="{color}" stroke="#1f2937" stroke-width="0.8"/>')
        elements.append(f'<text x="{cx+8:.1f}" y="{cy+4:.1f}" font-family="Arial, sans-serif" font-size="10" fill="#111827">{label}</text>')

    x_label = f'PC1 ({points["pca1_explained"].iloc[0]*100:.1f}%)'
    y_label = f'PC2 ({points["pca2_explained"].iloc[0]*100:.1f}%)'
    elements.append(f'<text x="{margin_left+plot_w/2}" y="{height-25}" text-anchor="middle" font-family="Arial, sans-serif" font-size="13">{x_label}</text>')
    elements.append(f'<text transform="translate(22 {margin_top+plot_h/2}) rotate(-90)" text-anchor="middle" font-family="Arial, sans-serif" font-size="13">{y_label}</text>')

    legend_x = margin_left + plot_w + 30
    legend_y = margin_top + 20
    elements.append(f'<text x="{legend_x}" y="{legend_y}" font-family="Arial, sans-serif" font-size="13" font-weight="700">Subclasses</text>')
    for idx, subclass in enumerate(subclasses):
        row = points[points["subclass_id"] == subclass].iloc[0]
        yy = legend_y + 24 + idx * 24
        elements.append(f'<rect x="{legend_x}" y="{yy-10}" width="12" height="12" fill="{color_map[subclass]}"/>')
        elements.append(f'<text x="{legend_x+18}" y="{yy}" font-family="Arial, sans-serif" font-size="12">{html.escape(subclass)} {html.escape(row["subclass_name"])}</text>')

    elements.append("</svg>")
    path.write_text("\n".join(elements), encoding="utf-8-sig")


def svg_major_rule(path: Path, major: pd.DataFrame, summary: pd.DataFrame) -> None:
    width, height = 780, 420
    margin_left, margin_right, margin_top, margin_bottom = 75, 35, 55, 70
    plot_w = width - margin_left - margin_right
    plot_h = height - margin_top - margin_bottom
    x = np.arange(len(major))
    y = major.sort_values("log_ratio_pbba_to_k")["log_ratio_pbba_to_k"].to_numpy()
    ordered = major.sort_values("log_ratio_pbba_to_k").reset_index(drop=True)
    y_min, y_max = float(y.min()), float(y.max())
    pad = 0.12 * (y_max - y_min)
    y_min -= pad
    y_max += pad
    threshold = float(summary.loc[summary["model"] == "log_ratio_threshold_in_sample", "threshold"].iloc[0])

    def sx(index: int) -> float:
        return margin_left + index / max(1, len(x) - 1) * plot_w

    def sy(value: float) -> float:
        return margin_top + (y_max - value) / (y_max - y_min) * plot_h

    color_map = {TYPE_HIGH_K: "#2f6f9f", TYPE_LEAD_BARIUM: "#c66b2b"}
    elements = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<rect width="100%" height="100%" fill="#ffffff"/>',
        f'<text x="{width/2}" y="28" text-anchor="middle" font-family="Arial, sans-serif" font-size="18" font-weight="700">Major-type log-ratio rule</text>',
        f'<line x1="{margin_left}" y1="{margin_top+plot_h}" x2="{margin_left+plot_w}" y2="{margin_top+plot_h}" stroke="#222" stroke-width="1"/>',
        f'<line x1="{margin_left}" y1="{margin_top}" x2="{margin_left}" y2="{margin_top+plot_h}" stroke="#222" stroke-width="1"/>',
        f'<line x1="{margin_left}" y1="{sy(threshold):.1f}" x2="{margin_left+plot_w}" y2="{sy(threshold):.1f}" stroke="#b91c1c" stroke-width="1.5" stroke-dasharray="6 4"/>',
        f'<text x="{margin_left+plot_w-5}" y="{sy(threshold)-6:.1f}" text-anchor="end" font-family="Arial, sans-serif" font-size="12" fill="#b91c1c">threshold={threshold:.3f}</text>',
    ]
    for idx, row in ordered.iterrows():
        cx = sx(idx)
        cy = sy(float(row["log_ratio_pbba_to_k"]))
        color = color_map[row["glass_type"]]
        elements.append(f'<circle cx="{cx:.1f}" cy="{cy:.1f}" r="5.5" fill="{color}" stroke="#111827" stroke-width="0.6"/>')
        elements.append(f'<text x="{cx:.1f}" y="{height-48}" text-anchor="middle" font-family="Arial, sans-serif" font-size="8" transform="rotate(-65 {cx:.1f} {height-48})">{html.escape(str(row["artifact_id"]))}</text>')
    elements.append(f'<text x="{margin_left+plot_w/2}" y="{height-18}" text-anchor="middle" font-family="Arial, sans-serif" font-size="13">Artifacts sorted by log((PbO+BaO+eps)/(K2O+eps))</text>')
    elements.append(f'<text transform="translate(24 {margin_top+plot_h/2}) rotate(-90)" text-anchor="middle" font-family="Arial, sans-serif" font-size="13">log-ratio</text>')
    elements.append(f'<rect x="{width-150}" y="58" width="12" height="12" fill="{color_map[TYPE_HIGH_K]}"/><text x="{width-132}" y="69" font-family="Arial, sans-serif" font-size="12">High-K</text>')
    elements.append(f'<rect x="{width-150}" y="80" width="12" height="12" fill="{color_map[TYPE_LEAD_BARIUM]}"/><text x="{width-132}" y="91" font-family="Arial, sans-serif" font-size="12">Lead-Ba</text>')
    elements.append("</svg>")
    path.write_text("\n".join(elements), encoding="utf-8-sig")


def write_summary_report(
    path: Path,
    artifacts: pd.DataFrame,
    major_summary: pd.DataFrame,
    subclass_centers: pd.DataFrame,
    raw_comparison: pd.DataFrame,
    pseudocount: pd.DataFrame,
    feature_sens: pd.DataFrame,
) -> None:
    major_acc = major_summary.loc[major_summary["model"] == "log_ratio_threshold_leave_one_out", "accuracy"].iloc[0]
    centroid_acc = major_summary.loc[major_summary["model"] == "clr_nearest_centroid_leave_one_out", "accuracy"].iloc[0]
    lines = [
        "# Question 2 Results Summary",
        "",
        "## Data",
        "",
        f"- Corrected artifact records: {len(artifacts)}.",
        f"- High-potassium records: {(artifacts['glass_type'] == TYPE_HIGH_K).sum()}.",
        f"- Lead-barium records: {(artifacts['glass_type'] == TYPE_LEAD_BARIUM).sum()}.",
        "",
        "## Major-Type Classification",
        "",
        f"- Leave-one-out accuracy of the log-ratio threshold rule: {major_acc:.3f}.",
        f"- Leave-one-out accuracy of the CLR nearest-centroid model: {centroid_acc:.3f}.",
        "- The main rule is based on log((PbO + BaO + eps) / (K2O + eps)).",
        "",
        "## Subclasses",
        "",
    ]
    for _, row in subclass_centers.iterrows():
        lines.append(
            f"- {row['glass_type']} {row['subclass_id']} {row['subclass_name']}: "
            f"n={int(row['n'])}, artifacts={row['artifact_ids']}."
        )
    lines += [
        "",
        "## Sensitivity",
        "",
        "- Raw-valid versus weathering-corrected adjusted Rand indices:",
    ]
    for _, row in raw_comparison.iterrows():
        lines.append(f"  - {row['glass_type']}: ARI={row['adjusted_rand_index']:.3f}.")
    lines.append("- Pseudocount sensitivity adjusted Rand indices:")
    for _, row in pseudocount.iterrows():
        lines.append(f"  - {row['glass_type']}, eps={row['pseudocount']}: ARI={row['adjusted_rand_index_vs_0_10']:.3f}.")
    lines.append("- Feature-deletion sensitivity adjusted Rand indices:")
    for _, row in feature_sens.iterrows():
        lines.append(f"  - {row['glass_type']}, drop {row['dropped_feature']}: ARI={row['adjusted_rand_index']:.3f}.")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8-sig")


def write_csv(frame: pd.DataFrame, path: Path) -> None:
    frame.to_csv(path, index=False, encoding="utf-8-sig", quoting=csv.QUOTE_MINIMAL)


def main() -> None:
    rng = np.random.default_rng(SEED)
    root = project_root()
    dirs = ensure_directories(root)

    workbook = find_source_workbook(root)
    _, raw_samples = read_workbook(workbook)
    predictions = read_preweathering_predictions(root)

    corrected_samples = build_corrected_samples(raw_samples, predictions)
    corrected_artifacts = aggregate_artifacts(corrected_samples)
    raw_artifacts = build_raw_artifact_data(raw_samples)

    write_csv(corrected_samples, dirs["processed"] / "corrected-sampling-points.csv")
    write_csv(corrected_artifacts, dirs["processed"] / "corrected-artifact-compositions.csv")
    write_csv(raw_artifacts, dirs["processed"] / "raw-artifact-compositions.csv")

    summary = component_summary(corrected_artifacts)
    major_results, major_summary = major_classification(corrected_artifacts)
    write_csv(summary, dirs["results"] / "type-component-summary.csv")
    write_csv(major_results, dirs["results"] / "major-type-classification.csv")
    write_csv(major_summary, dirs["results"] / "major-type-classification-summary.csv")

    high_k_assignments, high_k_centers, high_k_metrics, high_k_history, high_k_pca = subclass_analysis(
        corrected_artifacts, TYPE_HIGH_K, HIGH_K_FEATURES, 3, rng
    )
    lead_assignments, lead_centers, lead_metrics, lead_history, lead_pca = subclass_analysis(
        corrected_artifacts, TYPE_LEAD_BARIUM, LEAD_BARIUM_FEATURES, 4, rng
    )

    assignments = pd.concat([high_k_assignments, lead_assignments], ignore_index=True)
    centers = pd.concat([high_k_centers, lead_centers], ignore_index=True)
    metrics = pd.concat([
        high_k_metrics.assign(glass_type=TYPE_HIGH_K),
        lead_metrics.assign(glass_type=TYPE_LEAD_BARIUM),
    ], ignore_index=True)
    histories = pd.DataFrame([
        {"glass_type": TYPE_HIGH_K, **entry} for entry in high_k_history
    ] + [
        {"glass_type": TYPE_LEAD_BARIUM, **entry} for entry in lead_history
    ])
    pca_points = pd.concat([high_k_pca, lead_pca], ignore_index=True)

    write_csv(assignments, dirs["results"] / "subclass-assignments.csv")
    write_csv(centers, dirs["results"] / "subclass-centers.csv")
    write_csv(metrics, dirs["results"] / "cluster-k-selection.csv")
    write_csv(histories, dirs["results"] / "ward-merge-history.csv")
    write_csv(pca_points, dirs["results"] / "pca-subclass-coordinates.csv")

    raw_comparison = compare_raw_corrected(corrected_artifacts, raw_artifacts, assignments, rng)
    pseudo_sens = pseudocount_sensitivity(corrected_artifacts, assignments, rng)
    feature_sens = feature_sensitivity(corrected_artifacts, assignments, rng)
    boot_stability = bootstrap_stability(corrected_artifacts, assignments, rng)
    write_csv(raw_comparison, dirs["results"] / "raw-vs-corrected-sensitivity.csv")
    write_csv(pseudo_sens, dirs["results"] / "pseudocount-sensitivity.csv")
    write_csv(feature_sens, dirs["results"] / "feature-sensitivity.csv")
    write_csv(boot_stability, dirs["results"] / "bootstrap-subclass-stability.csv")

    svg_major_rule(dirs["figures"] / "major-type-log-ratio.svg", major_results, major_summary)
    svg_scatter(dirs["figures"] / "high-k-subclass-pca.svg", high_k_pca, "High-potassium subclass PCA")
    svg_scatter(dirs["figures"] / "lead-barium-subclass-pca.svg", lead_pca, "Lead-barium subclass PCA")

    write_summary_report(
        dirs["docs"] / "question-2-results-summary.md",
        corrected_artifacts,
        major_summary,
        centers,
        raw_comparison,
        pseudo_sens,
        feature_sens,
    )

    print("Question 2 analysis complete.")
    print(f"Corrected artifacts: {len(corrected_artifacts)}")
    print(f"Results: {dirs['results']}")
    print(f"Figures: {dirs['figures']}")


if __name__ == "__main__":
    main()
